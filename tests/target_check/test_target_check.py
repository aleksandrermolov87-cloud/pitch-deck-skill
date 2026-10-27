# Тест проверки PPTX в целевом приложении (scripts/pptx_target_check.py) и её места в сводном отчёте (check_report.py).
# Синтетические фикстуры: PPTX собирается python-pptx, «PDF из приложения» пишется вручную (адаптер fake), замер
# HTML-двойника подаётся файлом. Настоящие PowerPoint и LibreOffice тест не запускает: их отсутствие подставляется
# через PITCH_DECK_POWERPOINT_APP и PITCH_DECK_SOFFICE. Нужен poppler (pdftotext, pdffonts, pdfimages): без него — пропуск.
#   $SKILL/.venv/bin/python -I tests/target_check/test_target_check.py [папка]   (по умолчанию $SKILL/.build/target-check.XXXX)
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

SKILL = Path(__file__).resolve().parents[2]
PY = sys.executable
TC = SKILL / 'scripts' / 'pptx_target_check.py'
fails = []


def ok(cond, msg):
    print(('  ✓ ' if cond else '  ✗ ') + msg)
    if not cond:
        fails.append(msg)


def run(*args, env=None, cwd=None):
    e = dict(os.environ, PITCH_DECK_POWERPOINT_APP='/nonexistent/PowerPoint.app', PITCH_DECK_SOFFICE='/nonexistent/soffice', **(env or {}))
    p = subprocess.run([PY, '-I', *map(str, args)], capture_output=True, text=True, env=e, cwd=cwd)
    return p.returncode, p.stdout + p.stderr


# ───────── фикстуры ─────────
def make_pptx(path, image=True, text2='Repair costs rise'):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from PIL import Image
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, (txt, box, size) in enumerate((('Payouts grew to 154 bn', (1, 1, 6, 1), 28), (text2, (1, 1, 4, 0.6), 20))):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(*(Inches(v) for v in box))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        r = tf.paragraphs[0].add_run()
        r.text, r.font.name, r.font.size = txt, 'Styrene A LC', Pt(size)
        if i == 1 and image:
            b = io.BytesIO()
            Image.new('RGB', (40, 30), (200, 30, 30)).save(b, 'PNG')
            b.seek(0)
            s.shapes.add_picture(b, Inches(6), Inches(2), Inches(3), Inches(2))
    prs.save(path)


def make_pdf(path, pages, font):
    """pages: [(список (x, y_top, size, текст) в pt от верхнего левого угла, рисовать ли картинку)]."""
    objs = []
    add = lambda b: objs.append(b) or len(objs)
    add(b'')                                    # 1 каталог
    add(b'')                                    # 2 страницы
    f = add(f'<</Type/Font/Subtype/Type1/BaseFont/{font}/Encoding/WinAnsiEncoding>>'.encode())
    raw = bytes([200, 30, 30] * 4)
    im = add(b'<</Type/XObject/Subtype/Image/Width 2/Height 2/ColorSpace/DeviceRGB/BitsPerComponent 8/Length 12>>stream\n' + raw + b'\nendstream')
    kids = []
    for words, image in pages:
        ops = ''.join(f'BT /F1 {sz} Tf {x} {540 - y - sz * 0.8:.1f} Td ({t}) Tj ET\n' for x, y, sz, t in words)
        if image:
            ops += 'q 216 0 0 144 432 252 cm /Im1 Do Q\n'
        c = add(f'<</Length {len(ops)}>>stream\n{ops}endstream'.encode())
        kids.append(add(f'<</Type/Page/Parent 2 0 R/MediaBox[0 0 960 540]/Resources<</Font<</F1 {f} 0 R>>/XObject<</Im1 {im} 0 R>>>>/Contents {c} 0 R>>'.encode()))
    objs[0] = b'<</Type/Catalog/Pages 2 0 R>>'
    objs[1] = f'<</Type/Pages/Kids[{" ".join(f"{k} 0 R" for k in kids)}]/Count {len(kids)}>>'.encode()
    out, offs = bytearray(b'%PDF-1.4\n'), []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f'{i} 0 obj\n'.encode() + o + b'\nendobj\n'
    x = len(out)
    out += f'xref\n0 {len(objs) + 1}\n0000000000 65535 f \n'.encode() + b''.join(f'{o:010d} 00000 n \n'.encode() for o in offs)
    out += f'trailer\n<</Size {len(objs) + 1}/Root 1 0 R>>\nstartxref\n{x}\n%%EOF\n'.encode()
    Path(path).write_bytes(out)


def project(root, name='T'):
    src = root / f'src-{name}'
    (src / 'out').mkdir(parents=True)
    (src / 'data').mkdir()
    (src / 'deck.json').write_text(json.dumps({'name': name, 'title': name}), encoding='utf-8')
    make_pptx(root / f'{name}.pptx')
    layout = root / 'layout.json'
    layout.write_text(json.dumps({'font': True, 'items': [{'slide': 1, 'text': 'Payouts grew to 154 bn', 'lines': 1},
                                                          {'slide': 2, 'text': 'Repair costs rise', 'lines': 1}]}), encoding='utf-8')
    return src, layout


def state(src):
    return json.load(open(src / 'out' / 'target_check.json', encoding='utf-8'))


def main():
    if not all(shutil.which(t) for t in ('pdfinfo', 'pdftotext', 'pdffonts', 'pdfimages')):
        print('пропуск: нет poppler (brew install poppler) — разбор PDF не проверяется')
        return 0
    (SKILL / '.build').mkdir(exist_ok=True)
    root = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.mkdtemp(prefix='target-check.', dir=SKILL / '.build'))
    root.mkdir(parents=True, exist_ok=True)
    src, layout = project(root)
    good, bad = root / 'good.pdf', root / 'bad.pdf'
    make_pdf(good, [([(72, 72, 28, 'Payouts grew to 154 bn')], False), ([(72, 72, 20, 'Repair costs rise')], True)], 'StyreneALC-Regular')
    make_pdf(bad, [([(72, 72, 28, 'Payouts grew'), (72, 108, 28, 'to 154'), (72, 144, 28, 'bn')], False),
                   ([(72, 72, 20, 'Repair costs rise')], False)], 'ArialMT')

    print('› нет приложения: «не выполнялось»')
    rc, out = run(TC, 'status', src)
    ok(rc == 0 and 'PPTX в целевом приложении: не выполнялось' in out, 'status без прогонов: не выполнялось')
    ok(not (src / 'out' / 'target_check.json').exists(), 'status без прогонов файл не пишет')
    rc, out = run(TC, 'run', src)
    st = state(src)
    ok(rc == 0 and st['runs']['dark']['status'] == 'not_run' and 'нет ни PowerPoint, ни LibreOffice' in st['runs']['dark']['reason'],
       'run --app auto без PowerPoint и LibreOffice: не выполнялось, причина записана')
    rc, out = run(TC, 'run', src, '--app', 'powerpoint')
    ok('PowerPoint не установлен' in out or 'не macOS' in out, 'run --app powerpoint без приложения: причина в выводе')
    ok(st['summary']['status'] == 'not_run', 'итог: не выполнялось')

    print('› fake-адаптер, PDF без расхождений: «пройдено»')
    rc, out = run(TC, 'run', src, '--app', 'fake', '--fake-pdf', good, '--html-layout', layout)
    st = state(src)
    r = st['runs']['dark']
    ok(rc == 0 and r['status'] == 'passed', f'run: пройдено (получено {r["status"]}: {[s["issues"] for s in r["slides"]]})')
    ok(all(s['status'] == 'passed' for s in r['slides']) and len(r['slides']) == 2, 'оба слайда: пройдено')
    ok('StyreneALC-Regular' in r['pdf_fonts'] and 'Styrene A LC' in r['fonts'], 'шрифты PDF и версии установленных записаны')
    ok((src / 'out' / 'target' / 'dark' / 'png' / 's2.png').is_file() or not shutil.which('pdftoppm'), 'PNG страниц в out/target/dark/png')
    ok(r['app'] == 'тестовый адаптер', 'тестовый адаптер назван тестовым')

    print('› fake-адаптер, PDF с дефектами: «ошибки»')
    rc, out = run(TC, 'run', src, '--app', 'fake', '--fake-pdf', bad, '--html-layout', layout)
    st = state(src)
    r = st['runs']['dark']
    iss = {s['n']: ' | '.join(i['text'] for i in s['issues']) for s in r['slides']}
    ok(rc == 1 and r['status'] == 'failed', 'run: ошибки, код 1')
    ok('не помещается в бокс (вниз' in iss[1], f'слайд 1: переполнение вниз ({iss[1][:120]})')
    ok('строк 3, в HTML-двойнике 1' in iss[1], 'слайд 1: переносы против HTML-двойника')
    ok('подмена шрифта: в PPTX «Styrene A LC», в PDF ArialMT' in iss[1], 'слайд 1: подмена шрифта')
    ok('изображения не отрисованы: в PPTX 1, в PDF 0' in iss[2], 'слайд 2: картинка не отрисована')

    print('› за холстом')
    off = root / 'off.pdf'
    make_pdf(off, [([(72, 72, 28, 'Payouts grew to 154 bn')], False), ([(900, 72, 20, 'Repair costs rise')], True)], 'StyreneALC-Regular')
    run(TC, 'run', src, '--app', 'fake', '--fake-pdf', off, '--html-layout', layout)
    r = state(src)['runs']['dark']
    t2 = ' | '.join(i['text'] for i in r['slides'][1]['issues'])
    ok('за холст' in t2 and r['slides'][0]['status'] == 'passed', f'слайд 2: текст за холстом ({t2[:100]})')

    print('› ручной просмотр и устаревание')
    rc, out = run(TC, 'manual', src, '--theme', 'dark', '--slides', '1,2', '--by', 'тест')
    ok('слайд 2' in out and state(src)['summary']['status'] == 'failed', 'ручной просмотр не отменяет найденную ошибку')
    make_pptx(root / 'T.pptx', text2='Repair costs rise again')       # пересборка PPTX
    rc, out = run(TC, 'status', src)
    S = state(src)['summary']
    ok(S['status'] == 'not_run' and 'прошлой сборке' in out, 'после пересборки прогон и ручная запись устарели: не выполнялось')
    rc, out = run(TC, 'manual', src, '--theme', 'dark', '--slides', '1', '--by', 'тест')
    ok(state(src)['summary']['status'] == 'partial', 'ручной просмотр 1 из 2 слайдов: частично')
    rc, out = run(TC, 'manual', src, '--theme', 'dark', '--slides', '1-2', '--by', 'тест', '--notes', 'шрифт на месте')
    S = state(src)['summary']
    ok(S['status'] == 'passed' and S['slides']['2']['dark']['manual'] == 'passed', 'ручной просмотр всех слайдов: пройдено (вручную)')
    rc, out = run(TC, 'manual', src, '--theme', 'dark', '--slides', '2', '--bad', '2')
    ok(rc == 1 and state(src)['summary']['status'] == 'failed', 'ручной просмотр с проблемой: ошибки, код 1')
    rc, out = run(TC, 'manual', src, '--theme', 'light', '--slides', '1')
    ok(rc != 0 and 'нет PPTX темы' in out, 'ручная запись для несобранной темы — отказ')
    rc, out = run(TC, 'manual', src, '--theme', 'dark', '--slides', '5')
    ok(rc != 0 and '--slides' in out, 'номер слайда вне деки — отказ')

    print('› запись до 1.8: out/powerpoint.json')
    src2, _ = project(root, 'L')
    (src2 / 'out' / 'powerpoint.json').write_text(json.dumps({'viewed': [{'theme': 'dark', 'slides': [1], 'by': 'агент'}]}), encoding='utf-8')
    rc, out = run(TC, 'status', src2, '--json', src2 / 'out' / 's.json')
    S = json.load(open(src2 / 'out' / 's.json', encoding='utf-8'))
    ok(S['status'] == 'partial' and S['slides']['1']['dark']['manual'] == 'passed', 'старый powerpoint.json читается как ручной просмотр')

    print('› сводный отчёт: PPTX отдельно от HTML, слайды, происхождение картинок')
    src3, _ = project(root, 'R')
    o = src3 / 'out'
    json.dump([{'n': 1, 'fn': 'cover', 'archetype': 'cover', 'title': 'Обложка', 'lib': [], 'custom': None, 'texts': [], 'notes': ''},
               {'n': 2, 'fn': 'trend', 'archetype': 'bars-trend', 'title': 'Сборы растут', 'lib': ['barsTrend'], 'custom': None, 'texts': [], 'notes': ''}],
              open(o / 'slides.json', 'w', encoding='utf-8'), ensure_ascii=False)
    json.dump([], open(o / 'facts.json', 'w'))
    (o / 'render.json').write_text('{"slides":2,"font":true,"overflow":[]}')
    d3 = src3 / 'img' / '3d'
    d3.mkdir(parents=True)
    for f in ('cover.png', 'cover_check.png', 'cover.json', 'sources.json', 'step.png', 'demo_obj.png', 'loose.png', 'user.png'):
        (d3 / f).write_bytes(b'x')
    json.dump({'provider_choice': {'provider': 'magicpatterns', 'chosen_by_user': True, 'when': '2026-10-27T10:00:00+03:00', 'quota': None},
               'note': 'тестовая запись',
               'images': [
                   {'file': 'img/3d/cover.jpg', 'slide': 1, 'provider': 'magicpatterns', 'tool': 'generate_image', 'prompt_sha': 'ab',
                    'source': 'https://cdn.magicpatterns.com/patterns/x.jpg', 'created': '2026-10-27T10:05:00+03:00', 'data_sent': 'generalized'},
                   {'file': 'img/3d/step.png', 'slide': 2, 'provider': 'codex-imagegen', 'tool': 'image_gen', 'prompt_sha': 'cd',
                    'source': None, 'created': '2026-10-27T10:06:00+03:00', 'data_sent': 'generalized'},
                   {'file': 'img/3d/user.png', 'slide': 2, 'provider': 'user', 'tool': None, 'source': None, 'created': None, 'data_sent': 'none'},
                   {'file': 'img/3d/demo_obj.png', 'slide': 1, 'provider': 'magicpatterns', 'tool': 'generate_image', 'prompt_sha': None,
                    'source': 'https://cdn.magicpatterns.com/patterns/d.jpg', 'created': None, 'data_sent': 'none', 'demo': True}]},
              open(src3 / 'data' / 'visuals.json', 'w', encoding='utf-8'), ensure_ascii=False)
    cmpd = root / 'cmp'
    cmpd.mkdir()
    json.dump({'kind': 'контрольная повторная сборка', 'A': {'label': 'old'}, 'B': {'label': 'new'}, 'slides': [], 'reasons_not_controlled': [],
               'counts': {'identical': 2, 'changed': 0, 'missing_in_A': 0, 'missing_in_B': 0}, 'content': None,
               'content_compare': {'comparable': False, 'reason': 'A: манифест старого формата'}}, open(cmpd / 'report.json', 'w', encoding='utf-8'), ensure_ascii=False)
    run(SKILL / 'scripts' / 'check_report.py', '.', '--compare', cmpd, cwd=src3)
    md = (o / 'checks.md').read_text(encoding='utf-8')
    ok('PPTX в целевом приложении: не выполнялось' in md and 'проверен только HTML-двойник' in md, 'шапка: HTML-превью отдельно, PPTX не назван проверенным')
    ok('### 4. Просмотр в PowerPoint и проверка PPTX в целевом приложении: не выполнялась' in md, 'раздел 4: не выполнялась')
    ok('| Слайд | Заголовок | Компонент | Почему выбран | Входные данные | Источник | Ограничения | Проверка |' in md, 'таблица по слайдам: все столбцы')
    ok('«Как меняется во времени?»' in md and 'HTML: тёмная ✓; PPTX: не выполнялось' in md, 'слайд 2: причина выбора из каталога и степень проверки')
    ok('бриф не размечен' in md, 'нет requirements.json — так и написано')
    vis = md[md.index('### Происхождение изображений'):]
    mp, cx, us, dm = (vis.index(h) for h in ('**MagicPatterns', '**Codex imagegen', '**Картинки пользователя', '**Демо-объекты заготовки'))
    ok(mp < cx < us < dm, 'провайдеры в отдельных группах, демо-объекты отдельно и последними')
    ok('img/3d/cover.jpg |' in vis[mp:cx] and 'img/3d/step.png' in vis[cx:us] and 'img/3d/demo_obj.png' in vis[dm:], 'каждая картинка в группе своего провайдера')
    ok('Выбор пользователя: да' in vis and 'неизвестны' in vis and 'обобщённый промпт' in vis and 'cdn.magicpatterns.com' in vis, 'выбор, квота, data_sent, источник')
    ok('img/3d/loose.png' in vis and 'cover_check' not in vis and 'sources.json' not in vis and 'cover.png' not in vis.split('**Происхождение не записано:**')[1],
       'без записи — только loose.png; x.jpg ↔ x.png по имени, *_check.png и *.json не считаются')
    ok('img/3d/loose.png: происхождение не записано' in md, 'предупреждение «происхождение не записано» в разделе 5')
    ok('Содержание не удалось сравнить: A: манифест старого формата' in md, '--compare: причина несравнимого содержания')
    J = json.load(open(o / 'checks.json', encoding='utf-8'))
    ok(J['target']['status'] == 'not_run' and J['visuals']['unrecorded'] == ['img/3d/loose.png'] and len(J['slides']) == 2, 'checks.json: target, slides, visuals')

    print('› сводный отчёт с пройденной проверкой в приложении')
    shutil.copyfile(root / 'R.pptx', root / 'R.bak')
    run(TC, 'run', src3, '--app', 'fake', '--fake-pdf', good, '--html-layout', layout)
    run(SKILL / 'scripts' / 'check_report.py', '.', cwd=src3)
    md = (o / 'checks.md').read_text(encoding='utf-8')
    ok('PPTX в целевом приложении: пройдено' in md and 'тёмная тестовый адаптер: пройдено' in md, 'раздел 4 и таблица: пройдено в приложении по слайдам')

    if fails:
        print(f'\nтест НЕ пройден ({len(fails)}): {root}')
        return 1
    if len(sys.argv) < 2:
        shutil.rmtree(root, ignore_errors=True)
    print('\nтест пройден')
    return 0


if __name__ == '__main__':
    sys.exit(main())

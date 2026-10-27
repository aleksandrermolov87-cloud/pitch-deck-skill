# Сверка чисел на слайдах с источниками. facts.json пишет build.js через fact() из lib.js.
# Запуск: $SKILL/.venv/bin/python -I facts_check.py out/facts.json [--root .] [--xlsx model.xlsx] [--json out/facts_report.json]
# Типы источников: xlsx (ячейки, агрегаты value/sum/ratio/share_of), text (цитата в файле), user (сообщение пользователя),
# derived (из других фактов). Код выхода 1 при любом расхождении.
# Факт с auto: true — число, которое компонент нарисовал (высота столбца, точка линии), но не подписал. Он проверяется
# по источнику так же, как подписанный. shown (пишет build.js) — читается ли число на слайде: отчёт различает
# «влияет на график» и «можно прочитать». Что из брифа обязано читаться, проверяет brief_coverage.py, а не этот скрипт.
# shown_as (build.js): label — число читается вместе с периодом и единицей в одной подписи, value — читается только число
# (период на оси, единица в подзаголовке), graph — только нарисовано. Идентичность факта компонента: lib, path, series,
# period, unit, id. calc — запись расчёта (остаток stackedShare, сверка суммы с точностью) выводится в [INFO].
import argparse
import json
import os
import re
import sys
from collections import defaultdict

ap = argparse.ArgumentParser()
ap.add_argument('facts')
ap.add_argument('--root', default='.', help='папка, от которой считаются пути file в источниках text')
ap.add_argument('--xlsx', default=None, help='модель по умолчанию для фактов без поля file')
ap.add_argument('--html', default=None, help='out/deck.html: найти на слайдах числа, не зарегистрированные через fact()')
ap.add_argument('--allow', default='', help='через запятую числа, которые не требуют источника (например, 3,5)')
ap.add_argument('--json', default=None)
a = ap.parse_args()

F = json.load(open(a.facts))
books = {}


def book(path):
    if path not in books:
        import openpyxl
        books[path] = openpyxl.load_workbook(path, data_only=True)
    return books[path]


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def norm(s):
    return re.sub(r'\s+', ' ', s.replace(' ', ' ').replace(' ', ' ').replace(' ', ' ')).strip().lower()


def variants(v, digits):
    out = set()
    for d in {digits, 0, 1, 2}:
        s = f'{v:.{d}f}'
        out |= {s, s.replace('.', ','), f'{v:,.{d}f}'.replace(',', ' ').replace('.', ',')}
    return out


errors, notes, ok = [], [], 0
by_label = {f['label']: f for f in F}          # заранее: производный факт может стоять раньше своих слагаемых
for f in F:
    tol = 0.5 * 10 ** -f.get('digits', 0) + 1e-9
    kind = f.get('type', 'xlsx' if 'cells' in f else 'user')
    where = f"слайд {f['slide']}, «{f['label']}» = {f['value']}"
    try:
        if kind == 'xlsx':
            path = f.get('file') or a.xlsx
            if not path:
                errors.append(f'{where}: не указан xlsx (--xlsx или поле file)')
                continue
            ws = book(os.path.join(a.root, path) if f.get('file') else path)[f['sheet']]
            vals = [num(ws[c].value) for c in f['cells']]
            agg = f.get('agg', 'value')
            if agg in ('value', 'sum'):
                got = sum(vals)
            elif agg == 'ratio':
                got = vals[0] / vals[1]
            elif agg == 'share_of':
                got = sum(vals) / num(ws[f['den']].value)
            else:
                errors.append(f'{where}: неизвестный agg {agg}')
                continue
            got /= f.get('scale', 1)
            if abs(got - f['value']) <= tol:
                ok += 1
            else:
                errors.append(f"{where}: в {f['sheet']}!{','.join(f['cells'])} ({agg}) получается {got:.4f}")
        elif kind == 'text':
            text = norm(open(os.path.join(a.root, f['file']), encoding='utf-8').read())
            q = norm(f.get('quote', ''))
            if not q or q not in text:
                errors.append(f"{where}: цитата не найдена в {f['file']}: «{f.get('quote', '')[:60]}»")
            elif not any(v in q for v in variants(f['value'], f.get('digits', 0))):
                errors.append(f"{where}: числа нет в цитате «{f['quote'][:60]}»")
            else:
                ok += 1
        elif kind == 'user':
            notes.append(f"{where}: из сообщения пользователя {f.get('date', '')}: «{f.get('quote', '')[:60]}» (проверено вручную)")
            ok += 1
        elif kind == 'none':
            errors.append(f'{where}: число рисует график, но у него нет источника — добавьте src в данные компонента (item.src, data.src или opt.src)')
        elif kind == 'derived':
            src = [by_label.get(l) for l in f['from']]
            if None in src:
                errors.append(f"{where}: нет исходных фактов {f['from']}")
                continue
            v = [s['value'] for s in src]
            # rest — остаток до целого: whole − сумма частей (stackedShare, rest.type = computed)
            got = {'count': len(v), 'sum': sum(v), 'ratio': v[0] / v[1] if len(v) > 1 and v[1] else float('nan'), 'diff': v[0] - v[1] if len(v) > 1 else float('nan'),
                   'rest': f.get('whole', 100) - sum(v)}[f['op']] / f.get('scale', 1)
            if abs(got - f['value']) <= tol:
                ok += 1
            else:
                errors.append(f"{where}: {f['op']}({f['from']}) = {got:.4f}")
        else:
            errors.append(f'{where}: неизвестный тип источника {kind}')
    except (KeyError, ValueError, ZeroDivisionError, FileNotFoundError) as e:
        errors.append(f'{where}: {type(e).__name__}: {e}')

# одно и то же число под одной подписью на разных слайдах с разными значениями
seen = defaultdict(set)
for f in F:
    if not f.get('auto'):                      # подписи auto-фактов — путь в данных компонента, на разных слайдах они совпадают законно
        seen[f['label']].add((f['slide'], f['value']))
for label, s in seen.items():
    if len({v for _, v in s}) > 1:
        errors.append(f"«{label}»: разные значения на слайдах {sorted(s)}")
    elif len({n for n, _ in s}) > 1:
        notes.append(f"«{label}» повторяется на слайдах {sorted({n for n, _ in s})}: проверьте, нужен ли дубль")

# числа на слайдах без fact(): вписаны строкой, источник не проверен
unregistered = []
if a.html:
    import html as H
    doc = open(a.html, encoding='utf-8').read()
    allow = {x.strip() for x in a.allow.split(',') if x.strip()}
    # разрешённые тексты из deck.json: определения вроде «сегмент 0–5», числа в них не факты
    dj = os.path.join(a.root, 'deck.json')
    ok_texts = [(e['slide'], e['text']) for e in json.load(open(dj)).get('numbers_ok', [])] if os.path.isfile(dj) else []
    per_slide = defaultdict(set)
    for f in F:
        # модуль числа тоже: на слайде минус типографский («−9»), а токен выделяется без знака
        per_slide[f['slide']] |= {norm(v) for v in variants(abs(f['value']), f.get('digits', 0))}
    for m in re.finditer(r'<section class="slide" id="s(\d+)".*?</section>', doc, re.S):
        n = int(m.group(1))
        # текст по элементам: номер страницы «2 / 14» — отдельный элемент, числа соседних элементов не склеиваются
        for el in re.findall(r'<div class="t"[^>]*>(.*?)</div>', m.group(0), re.S):
            txt = H.unescape(re.sub(r'<[^>]+>', '', el)).strip()
            if re.fullmatch(r'\d+\s*/\s*\d+', txt) or any(sn == n and t in txt for sn, t in ok_texts):
                continue
            # даты «15 декабря», «3 мар», «12–18 янв» — не факты: убираем их до поиска чисел
            txt = re.sub(r'\b\d{1,2}(?:\s*[–-]\s*\d{1,2})?[\s\u00a0](?:янв|фев|мар|апр|ма[йя]|июн|июл|авг|сен|окт|ноя|дек)[а-я]*\.?', ' ', txt)
            for tok in re.findall(r'(?<![\w.,])\d[\d\u00a0\u202f ]*(?:[.,]\d+)?', txt):
                t = norm(tok).strip()
                if not t or t in allow or re.fullmatch(r'20\d\d|19\d\d', t) or re.fullmatch(r'\d{1,2}\.\d{1,2}(\.\d{2,4})?', t):
                    continue
                if t.replace(' ', '') not in {v.replace(' ', '') for v in per_slide[n]}:
                    unregistered.append(f'слайд {n}: «{t}» в «{txt[:40]}» вписано без fact()')

for e in errors:
    print('[ERR]', e)
for u in unregistered:
    print('[WARN]', u)
# числа, которые только рисуют график: по слайдам, чтобы было видно, что зритель их не прочтёт
graph_only = [f for f in F if f.get('shown') is False and (f.get('lib') or f.get('auto'))]
idle = [f for f in F if f.get('shown') is False and not (f.get('lib') or f.get('auto')) and f.get('type') != 'derived']
for f in idle:
    notes.append(f"слайд {f['slide']}: «{f['label']}» = {f['value']} зарегистрирован, но на слайде не читается и компонентом не нарисован (заметки или лишний fact())")
if graph_only:
    per = defaultdict(list)
    for f in graph_only:
        per[f['slide']].append(f"{f['label'].split(' [')[0]} = {f['value']}")
    for n_, xs in sorted(per.items()):
        notes.append(f"слайд {n_}: только в графике, без подписи ({len(xs)}): " + '; '.join(xs[:6]) + (' …' if len(xs) > 6 else ''))
for f in F:
    if f.get('calc'):
        notes.append(f"слайд {f['slide']}: расчёт «{f['label']}»: {f['calc']}")
for n in notes:
    print('[INFO]', n)
shown_n = sum(1 for f in F if f.get('shown'))
labeled = sum(1 for f in F if f.get('shown') and f.get('shown_as') == 'label')
value_only = sum(1 for f in F if f.get('shown') and f.get('shown_as') == 'value')
print(f'\nфактов {len(F)}, сошлось {ok}, расхождений {len(errors)}, чисел без источника {len(unregistered)}'
      + (f'; читаются на слайде {shown_n}, только в графике {len(graph_only)}' if any('shown' in f for f in F) else '')
      + (f'; подписаны с периодом и единицей {labeled}, только число {value_only}' if any('shown_as' in f for f in F) else ''))
if a.json:
    json.dump({'total': len(F), 'ok': ok, 'errors': errors, 'unregistered': unregistered, 'notes': notes,
               'shown': shown_n, 'graph_only': len(graph_only), 'labeled': labeled, 'value_only': value_only},
              open(a.json, 'w'), ensure_ascii=False, indent=1)
sys.exit(1 if errors or unregistered else 0)

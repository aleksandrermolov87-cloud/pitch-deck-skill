# Проверка готового PPTX в целевом приложении (reference/17-check-report.md, «PPTX в целевом приложении»). Профиль
# по желанию: для выпуска не обязателен, но без него отчёт не называет PPTX проверенным.
#   pptx_target_check.py run    <src> [--app auto|powerpoint|libreoffice|none] [--theme all|dark|light] [--timeout 240]
#   pptx_target_check.py manual <src> --theme dark --slides 1,3,5 [--bad 3] [--app powerpoint] [--by "…"] [--notes "…"]
#   pptx_target_check.py status <src> [--json файл]      итог по темам и слайдам с учётом пересборок
#   pptx_target_check.py env                             какие адаптеры и инструменты PDF есть в окружении
# Адаптеры по доступности: Microsoft PowerPoint на macOS (AppleScript, экспорт в PDF), LibreOffice
# (soffice --headless --convert-to pdf). Настройки системы скрипт не меняет: если macOS спросит разрешение
# на управление PowerPoint или PowerPoint попросит доступ к папке, решает пользователь; без ответа — «не выполнялось».
# По PDF (poppler: pdftotext, pdffonts, pdfimages, pdftoppm) в границах возможностей приложения проверяет:
# переполнение бокса и выход за холст, переносы строк против HTML-двойника (замер в Chrome), отрисованы ли
# изображения, подменён ли шрифт (с версиями установленных шрифтов). PNG страниц — out/target/<тема>/png.
# Пишет out/target_check.json. Статусы по слайду и в целом: «не выполнялось», «частично», «пройдено», «ошибки».
# Результат прогона и ручной записи привязан к sha256 PPTX: после пересборки он считается устаревшим.
# Код 1 у run и status, если найдены ошибки. Запуск: $SKILL/.venv/bin/python -I pptx_target_check.py …
import argparse
import datetime as dt
import hashlib
import json
import os
import plistlib
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
LABEL = {'not_run': 'не выполнялось', 'partial': 'частично', 'passed': 'пройдено', 'failed': 'ошибки'}
THEMES = {'dark': 'тёмная', 'light': 'светлая'}
PP_APP = os.environ.get('PITCH_DECK_POWERPOINT_APP', '/Applications/Microsoft PowerPoint.app')
EMU = 914400


def now():
    return dt.datetime.now().astimezone().isoformat(timespec='seconds')


def sha(p):
    h = hashlib.sha256()
    with open(p, 'rb') as f:
        for b in iter(lambda: f.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()


def sh(cmd, timeout=60, **kw):
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, **kw)
        return p.returncode, p.stdout, p.stderr
    except subprocess.TimeoutExpired:
        return -9, '', f'нет ответа за {timeout} с'
    except OSError as e:
        return -1, '', str(e)


# ───────── окружение: адаптеры и инструменты ─────────
def powerpoint():
    if sys.platform != 'darwin':
        return {'name': 'powerpoint', 'title': 'Microsoft PowerPoint', 'available': False, 'why': 'не macOS: адаптер PowerPoint работает только на macOS'}
    app = Path(PP_APP)
    if not (app / 'Contents' / 'Info.plist').is_file():
        return {'name': 'powerpoint', 'title': 'Microsoft PowerPoint', 'available': False, 'why': f'PowerPoint не установлен ({app})'}
    try:
        ver = plistlib.load(open(app / 'Contents' / 'Info.plist', 'rb')).get('CFBundleShortVersionString', '?')
    except Exception:
        ver = '?'
    return {'name': 'powerpoint', 'title': 'Microsoft PowerPoint', 'available': True, 'version': ver, 'path': str(app),
            'fonts': [str(app / 'Contents' / 'Resources' / 'DFonts')]}


def libreoffice():
    cand = [os.environ.get('PITCH_DECK_SOFFICE')] if 'PITCH_DECK_SOFFICE' in os.environ else \
        [shutil.which('soffice'), shutil.which('libreoffice'), '/Applications/LibreOffice.app/Contents/MacOS/soffice']
    exe = next((c for c in cand if c and os.access(c, os.X_OK)), None)
    if not exe:
        return {'name': 'libreoffice', 'title': 'LibreOffice', 'available': False, 'why': 'LibreOffice (soffice) не найден'}
    return {'name': 'libreoffice', 'title': 'LibreOffice', 'available': True, 'path': exe, 'version': None}


def lo_version(ad):
    rc, out, _ = sh([ad['path'], '--version'], timeout=60)
    m = re.search(r'LibreOffice\s+([\d.]+)', out)
    return m.group(1) if m else '?'


TOOLS = ('pdfinfo', 'pdftotext', 'pdffonts', 'pdfimages', 'pdftoppm')


def tools():
    return {t: shutil.which(t) for t in TOOLS}


def environment():
    return {'platform': sys.platform, 'adapters': [powerpoint(), libreoffice()], 'tools': tools(), 'chrome': chrome()}


def chrome():
    rc, out, _ = sh(['zsh', str(SKILL / 'scripts' / 'chrome.sh')], timeout=10)
    return out.strip() if rc == 0 and out.strip() else None


# ───────── проект: темы, PPTX, слайды ─────────
def project(src):
    deck = json.load(open(src / 'deck.json', encoding='utf-8')) if (src / 'deck.json').is_file() else {}
    name = deck.get('name', src.name)
    up = src.parent
    themes = {}
    for th, sfx in (('dark', '_темная'), ('dark', ''), ('light', '_светлая')):
        p = up / f'{name}{sfx}.pptx'
        if p.is_file() and th not in themes:
            themes[th] = p
    return name, themes


def html_of(src, theme):
    return src / ('out' if theme == 'dark' else 'out-light') / 'deck.html'


def png_of(src, theme, n):
    return src / ('out' if theme == 'dark' else 'out-light') / 'png' / f's{n}.png'


def pptx_model(path):
    """Слайды PPTX: размер, скрытость, текстовые боксы (доли холста), картинки, шрифты."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    slides = []
    for i, s in enumerate(prs.slides, 1):
        texts, pics = [], 0
        for sh_ in s.shapes:
            if sh_.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pics += 1
                try:
                    sh_.image.blob
                except Exception:
                    pics -= 1
                    texts.append({'broken_image': True})
            if getattr(sh_, 'has_text_frame', False) and sh_.has_text_frame and sh_.text_frame.text.strip():
                texts.append({'text': sh_.text_frame.text, 'box': [sh_.left / W, sh_.top / H, sh_.width / W, sh_.height / H]})
        bg = s._element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
        if bg is not None and bg.find(f'.//{A}blip') is not None:
            pics += 1
        fonts = sorted({e.get('typeface') for e in s._element.iter(f'{A}latin') if e.get('typeface') and not e.get('typeface').startswith('+')})
        slides.append({'n': i, 'hidden': s._element.get('show') == '0', 'texts': [t for t in texts if 'text' in t],
                       'broken_images': sum(1 for t in texts if t.get('broken_image')), 'images': pics, 'fonts': fonts})
    return {'w_in': W / EMU, 'h_in': H / EMU, 'slides': slides}


# ───────── нормализация текста ─────────
def toks(s):
    s = unicodedata.normalize('NFKC', s).replace('­', '').replace('−', '-').lower()
    out = []
    for w in s.split():
        w = w.strip('.,;:!?«»"„“”()[]—–-…\'’%‰₽$€£+±×*#/')   # знаки при числе PDF часто отдаёт отдельным словом
        if any(c.isalnum() for c in w):
            out.append(w)
    return out


# ───────── HTML-двойник: строки каждого текста (замер в Chrome) ─────────
MEASURE = ("<script>document.fonts.ready.then(()=>{const R=[];document.querySelectorAll('section.slide').forEach((s,si)=>{"
           "s.querySelectorAll('.t').forEach(t=>{const tops=[],rs=[],w=document.createTreeWalker(t,NodeFilter.SHOW_TEXT);"
           "while(w.nextNode()){const r=document.createRange();r.selectNodeContents(w.currentNode);rs.push(...r.getClientRects());}"
           "for(const q of rs){if(q.width<1||q.height<1)continue;const c=(q.top+q.bottom)/2,l=tops.find(l=>c>=l[0]&&c<=l[1]);"
           "if(l){l[0]=Math.min(l[0],q.top);l[1]=Math.max(l[1],q.bottom);}else tops.push([q.top,q.bottom]);}"
           "R.push({slide:si+1,text:t.textContent,lines:tops.length});});});const o=document.createElement('pre');o.id='M';"
           "o.textContent=JSON.stringify({font:document.fonts.check('12px \"FONT\"'),items:R});document.body.appendChild(o);});</script></body>")


def html_layout(src, theme, cache_dir, given=None):
    """{slide: [(tokens, lines)]} или (None, причина)."""
    if given:
        data = json.load(open(given, encoding='utf-8'))
    else:
        html = html_of(src, theme)
        if not html.is_file():
            return None, f'нет {html.relative_to(src)}: переносы с HTML-двойником не сравнивались'
        key = sha(html)
        cache = cache_dir / 'html_layout.json'
        data = json.load(open(cache, encoding='utf-8')) if cache.is_file() else {}
        if data.get('deck_sha') != key:
            exe = chrome()
            if not exe:
                return None, 'нет Chrome: переносы с HTML-двойником не сравнивались'
            font = ''
            if (src / 'theme.json').is_file():
                font = json.load(open(src / 'theme.json', encoding='utf-8')).get('font', {}).get('regular', '')
            with tempfile.TemporaryDirectory(prefix='pitch-deck-target-') as tmp:
                m = Path(tmp) / 'm.html'
                m.write_text(html.read_text(encoding='utf-8').replace('</body>', MEASURE.replace('FONT', font.replace('"', '')), 1), encoding='utf-8')
                rc, out, err = sh(['perl', '-e', 'alarm 60; exec @ARGV', exe, '--headless=new', '--disable-gpu', '--use-mock-keychain',
                                   '--password-store=basic', '--no-first-run', f'--user-data-dir={tmp}/p', '--window-size=1920,1080',
                                   '--virtual-time-budget=3000', '--dump-dom', m.as_uri()], timeout=90)
            mm = re.search(r'<pre id="M">(.*?)</pre>', out, re.S)
            if not mm:
                return None, 'замер HTML-двойника в Chrome не удался: переносы не сравнивались'
            import html as H
            data = json.loads(H.unescape(mm.group(1)))
            data['deck_sha'] = key
            cache_dir.mkdir(parents=True, exist_ok=True)
            json.dump(data, open(cache, 'w', encoding='utf-8'), ensure_ascii=False)
    lay = {}
    for it in data.get('items', []):
        lay.setdefault(int(it['slide']), []).append((toks(it['text']), int(it['lines'])))
    note = None if data.get('font', True) else 'шрифт бренда не загрузился в Chrome: строки HTML-двойника посчитаны на запасном шрифте'
    return lay, note


# ───────── адаптеры: PPTX → PDF ─────────
OSA = '''on run argv
  set src to POSIX file (item 1 of argv)
  set dst to (POSIX file (item 2 of argv)) as string
  tell application id "com.microsoft.Powerpoint"
    open src
    set p to active presentation
    save p in dst as save as PDF
    close p saving no
  end tell
end run'''


def export_powerpoint(ad, pptx, pdf, timeout):
    was = sh(['pgrep', '-xq', 'Microsoft PowerPoint'], timeout=5)[0] == 0
    rc, out, err = sh(['osascript', '-', str(pptx), str(pdf)], timeout=timeout, input=OSA)
    if not was:
        sh(['osascript', '-e', 'tell application id "com.microsoft.Powerpoint" to quit saving no'], timeout=30)
    if rc == 0 and pdf.is_file():
        return None
    if '-1743' in err or 'Not authorized' in err or 'не разрешено' in err:
        return ('macOS не дала разрешения управлять PowerPoint (Automation). Навык настройки не меняет: разрешите в окне macOS '
                'или в «Конфиденциальность → Автоматизация» и повторите run')
    if rc == -9:
        return (f'PowerPoint не ответил за {timeout} с: возможно, ждёт ответа в своём окне (разрешение доступа к папке, '
                'восстановление файла). Ответьте в окне PowerPoint и повторите run')
    return f'экспорт PowerPoint в PDF не удался: {(err or out).strip()[:200]}'


def export_libreoffice(ad, pptx, pdf, timeout):
    with tempfile.TemporaryDirectory(prefix='pitch-deck-lo-') as tmp:
        rc, out, err = sh([ad['path'], '--headless', '--norestore', '--nolockcheck', f'-env:UserInstallation={Path(tmp, "profile").as_uri()}',
                           '--convert-to', 'pdf', '--outdir', tmp, str(pptx)], timeout=timeout)
        got = Path(tmp) / (pptx.stem + '.pdf')
        if got.is_file():
            shutil.move(str(got), pdf)
            return None
    return f'LibreOffice не сконвертировал PPTX: {(err or out).strip()[:200] or f"код {rc}"}'


def export_fake(ad, pptx, pdf, timeout):
    shutil.copyfile(ad['pdf'], pdf)
    return None


# ───────── шрифты: версии установленных ─────────
def _sfnt_names(f, off):
    f.seek(off)
    hdr = f.read(12)
    if len(hdr) < 12:
        return {}
    n = struct.unpack('>H', hdr[4:6])[0]
    recs = f.read(16 * n)
    for i in range(n):
        tag, _, toff, tlen = struct.unpack('>4sIII', recs[16 * i:16 * i + 16])
        if tag != b'name':
            continue
        f.seek(toff)
        t = f.read(tlen)
        _, count, soff = struct.unpack('>HHH', t[:6])
        best = {}
        for k in range(count):
            pid, eid, lid, nid, ln, o = struct.unpack('>HHHHHH', t[6 + 12 * k:18 + 12 * k])
            if nid not in (1, 4, 5, 6, 16):
                continue
            raw = t[soff + o:soff + o + ln]
            s = raw.decode('utf-16-be', 'replace') if pid in (0, 3) else raw.decode('mac_roman', 'replace')
            rank = 0 if (pid == 3 and lid == 0x409) else 1 if pid == 3 else 2
            if nid not in best or rank < best[nid][0]:
                best[nid] = (rank, s.strip('\x00 '))
        return {k: v[1] for k, v in best.items()}
    return {}


def font_index(extra=()):
    dirs = [Path.home() / 'Library/Fonts', Path('/Library/Fonts'), Path('/System/Library/Fonts'), Path('/System/Library/Fonts/Supplemental'),
            Path('/usr/share/fonts'), Path('/usr/local/share/fonts'), Path.home() / '.fonts', Path.home() / '.local/share/fonts', *map(Path, extra)]
    idx, seen = [], 0
    for d in dirs:
        if not d.is_dir():
            continue
        for p in d.rglob('*'):
            if p.suffix.lower() not in ('.ttf', '.otf', '.ttc') or seen > 4000:
                continue
            seen += 1
            try:
                with open(p, 'rb') as f:
                    head = f.read(12)
                    offs = [0]
                    if head[:4] == b'ttcf':
                        num = struct.unpack('>I', head[8:12])[0]
                        offs = list(struct.unpack(f'>{min(num, 64)}I', f.read(4 * min(num, 64))))
                    for o in offs:
                        nm = _sfnt_names(f, o)
                        if nm:
                            idx.append({'file': str(p), 'family': nm.get(16) or nm.get(1, ''), 'full': nm.get(4, ''), 'ps': nm.get(6, ''),
                                        'version': re.sub(r'^Version\s*', '', nm.get(5, ''), flags=re.I).split(';')[0][:40]})
            except Exception:
                continue
    return idx


def fnorm(s):
    return re.sub(r'[^0-9a-zа-яё]', '', (s or '').lower())


def font_match(family, pdf_name):
    a, b = fnorm(family), fnorm(re.sub(r'^[A-Z]{6}\+', '', pdf_name or ''))
    base = fnorm(re.sub(r'^[A-Z]{6}\+', '', pdf_name or '').split('-')[0].split(',')[0])
    return bool(a) and (b.startswith(a) or (base and a.startswith(base) and len(base) >= 4))


def installed_versions(families, pdf_fonts, idx):
    out = {}
    for fam in families:
        hits = [x for x in idx if fnorm(x['family']) == fnorm(fam) or fnorm(x['full']).startswith(fnorm(fam))]
        out[fam] = sorted({f"{x['full'] or x['family']} {x['version']}".strip() for x in hits})[:4] or ['не установлен']
    for pf in pdf_fonts:
        ps = re.sub(r'^[A-Z]{6}\+', '', pf)
        hits = [x for x in idx if fnorm(x['ps']) == fnorm(ps)]
        if hits and ps not in out:
            out[ps] = sorted({f"{x['full'] or x['ps']} {x['version']}" for x in hits})[:2]
    return out


# ───────── анализ PDF ─────────
def pdf_pages(pdf, T):
    if not T['pdfinfo']:
        return None
    rc, out, _ = sh([T['pdfinfo'], str(pdf)])
    m = re.search(r'Pages:\s+(\d+)', out)
    s = re.search(r'Page size:\s+([\d.]+) x ([\d.]+)', out)
    return {'pages': int(m.group(1)) if m else 0, 'w': float(s.group(1)) if s else 0, 'h': float(s.group(2)) if s else 0}


def pdf_words(pdf, T):
    """{page: [{'t', x0, y0, x1, y1}]} в pt; None, если pdftotext нет."""
    if not T['pdftotext']:
        return None
    rc, out, _ = sh([T['pdftotext'], '-bbox', '-enc', 'UTF-8', str(pdf), '-'], timeout=120)
    if rc != 0:
        return None
    import html as H
    pages, cur = {}, None
    for line in out.splitlines():
        if '<page ' in line:
            cur = len(pages) + 1
            pages[cur] = []
        m = re.search(r'<word xMin="([-\d.]+)" yMin="([-\d.]+)" xMax="([-\d.]+)" yMax="([-\d.]+)">(.*?)</word>', line)
        if m and cur:
            for t in toks(H.unescape(m.group(5))) or []:
                pages[cur].append({'t': t, 'x0': float(m.group(1)), 'y0': float(m.group(2)), 'x1': float(m.group(3)), 'y1': float(m.group(4))})
    return pages


def pdf_fonts(pdf, T, page):
    if not T['pdffonts']:
        return None
    rc, out, _ = sh([T['pdffonts'], '-f', str(page), '-l', str(page), str(pdf)])
    fonts = []
    for line in out.splitlines()[2:]:
        parts = line.split()
        if len(parts) < 6:
            continue
        # name type… encoding emb sub uni object ID: справа налево фиксированные колонки
        emb, sub = parts[-5], parts[-4]
        fonts.append({'name': parts[0], 'pdf_name': re.sub(r'^[A-Z]{6}\+', '', parts[0]), 'embedded': emb == 'yes', 'subset': sub == 'yes'})
    return fonts


def pdf_images(pdf, T):
    if not T['pdfimages']:
        return None
    rc, out, _ = sh([T['pdfimages'], '-list', str(pdf)])
    cnt = {}
    for line in out.splitlines()[2:]:
        p = line.split()
        if len(p) > 2 and p[0].isdigit() and p[2] == 'image':
            cnt[int(p[0])] = cnt.get(int(p[0]), 0) + 1
    return cnt


def match_text(words, used, tk):
    """Лучшее совпадение последовательности слов: (число совпавших слов текста, индексы слов PDF)."""
    best = (0, [])
    for i, w in enumerate(words):
        if i in used or not (w['t'] == tk[0] or tk[0].startswith(w['t'])):
            continue
        got, k, n = [], i, 0
        for t in tk:
            if k >= len(words) or k in used:
                break
            if words[k]['t'] == t:
                got.append(k); k += 1; n += 1
                continue
            acc, idx, kk = words[k]['t'], [k], k + 1        # слово, разорванное переносом
            while t.startswith(acc) and acc != t and kk < len(words) and kk not in used and len(idx) < 4:
                acc += words[kk]['t']; idx.append(kk); kk += 1
            if acc != t:
                break
            got += idx; k = kk; n += 1
        if n > best[0]:
            best = (n, got)
    return best


def lines_of(ws):
    """Строки по вертикали: слово с центром внутри уже найденной строки — та же строка (разные кегли на одной базовой линии)."""
    out = []
    for w in sorted(ws, key=lambda w: w['y1'] - w['y0'], reverse=True):
        c = (w['y0'] + w['y1']) / 2
        if not any(a <= c <= b for a, b in out):
            out.append((w['y0'], w['y1']))
    return len(out)


def check_page(sl, words, page, fonts, imgs, lay):
    """Проблемы одного слайда: [(уровень, текст)], плюс набор невыполненных проверок."""
    iss, skipped = [], []
    PW, PH = page['w'], page['h']
    short = lambda s: (s[:48] + '…') if len(s) > 48 else s
    if words is None:
        skipped.append('текст и переполнения (нет pdftotext)')
    else:
        if sl['texts'] and not words:
            iss.append(('err', 'на странице нет текста: шрифт не отрисован или текст ушёл в кривые'))
        used = set()
        hl = {}
        for tk, n in (lay or []):
            hl.setdefault(' '.join(tk), []).append(n)
        for t in sorted(sl['texts'], key=lambda t: -len(toks(t['text']))):
            tk = toks(t['text'])
            if sum(len(x) for x in tk) < 2 or not words:
                continue
            n, got = match_text(words, used, tk)
            if n < max(1, -(-len(tk) * 6 // 10)):
                iss.append(('err', f'текст «{short(t["text"].strip())}» не найден в PDF: не отрисован или обрезан за холстом'))
                continue
            used.update(got)
            ws = [words[k] for k in got]
            if n < len(tk):
                iss.append(('warn', f'«{short(t["text"].strip())}»: в PDF найдено {n} из {len(tk)} слов'))
            x0, y0 = min(w['x0'] for w in ws), min(w['y0'] for w in ws)
            x1, y1 = max(w['x1'] for w in ws), max(w['y1'] for w in ws)
            lh = sorted(w['y1'] - w['y0'] for w in ws)[len(ws) // 2]
            tol = max(0.3 * lh, 1.5)
            bx, by, bw, bh = (t['box'][0] * PW, t['box'][1] * PH, t['box'][2] * PW, t['box'][3] * PH)
            if x0 < -0.5 or y0 < -0.5 or x1 > PW + 0.5 or y1 > PH + 0.5:
                iss.append(('err', f'«{short(t["text"].strip())}» выходит за холст'))
            over = [d for d, v in (('вниз', y1 - (by + bh)), ('вверх', by - y0), ('вправо', x1 - (bx + bw)), ('влево', bx - x0)) if v > tol]
            if over:
                iss.append(('err', f'«{short(t["text"].strip())}» не помещается в бокс ({", ".join(over)}): '
                                   f'текст {(y1 - y0) / 72:.2f}″ при боксе {bh / 72:.2f}″'))
            nl = lines_of(ws)
            hn = hl.get(' '.join(tk))
            if lay is not None and hn and nl != hn[0]:
                iss.append(('warn', f'«{short(t["text"].strip())}»: строк {nl}, в HTML-двойнике {hn[0]}'))
        if lay is None:
            skipped.append('переносы против HTML-двойника')
    if imgs is None:
        skipped.append('изображения (нет pdfimages)')
    else:
        got = imgs.get(page['n'], 0)
        if sl['broken_images']:
            iss.append(('err', f'в PPTX {sl["broken_images"]} изображений ссылаются на отсутствующий файл'))
        if sl['images'] and not got:
            iss.append(('err', f'изображения не отрисованы: в PPTX {sl["images"]}, в PDF 0'))
        elif got < sl['images']:
            iss.append(('warn', f'изображений в PDF {got} из {sl["images"]} (приложение могло объединить фон и картинки)'))
    if fonts is None:
        skipped.append('шрифты (нет pdffonts)')
    elif sl['texts']:
        for fam in sl['fonts']:
            if not any(font_match(fam, f['pdf_name']) for f in fonts):
                got = ', '.join(sorted({f['pdf_name'] for f in fonts})) or 'шрифтов нет'
                iss.append(('err', f'подмена шрифта: в PPTX «{fam}», в PDF {got}'))
    return iss, skipped


def visual_diff(a, b):
    try:
        import numpy as np
        from PIL import Image
        x = np.asarray(Image.open(a).convert('L').resize((320, 180)), dtype=float)
        y = np.asarray(Image.open(b).convert('L').resize((320, 180)), dtype=float)
        return round(float(abs(x - y).mean() / 255 * 100), 1)
    except Exception:
        return None


def analyze(src, theme, pptx, pdf, app, work, layout_file=None):
    T = tools()
    model = pptx_model(pptx)
    info = pdf_pages(pdf, T)
    run = {'theme': theme, 'pptx': pptx.name, 'pptx_sha': sha(pptx), 'app': app['title'], 'adapter': app['name'],
           'app_version': app.get('version'), 'when': now(), 'pdf': str(pdf.relative_to(src)), 'errors': [], 'limits': [], 'slides': []}
    if info is None:
        run['limits'].append('нет poppler (pdfinfo): PDF не разобран, проверено только, что приложение открыло файл')
        info = {'pages': 0, 'w': 0, 'h': 0}
    slides = model['slides']
    vis = [s for s in slides if not s['hidden']]
    if info['pages'] == len(slides):
        pages = {s['n']: i + 1 for i, s in enumerate(slides)}
    elif info['pages'] == len(vis):
        pages = {s['n']: i + 1 for i, s in enumerate(vis)}
    else:
        pages = {}
        if info['pages']:
            run['errors'].append(f'страниц в PDF {info["pages"]}, слайдов в PPTX {len(slides)} (видимых {len(vis)}): слайды не сопоставлены')
    if info['w'] and abs(info['w'] / info['h'] - model['w_in'] / model['h_in']) > 0.01:
        run['errors'].append(f'пропорции страницы {info["w"]:.0f}×{info["h"]:.0f} pt не совпали с холстом {model["w_in"]:.3f}×{model["h_in"]:.3f}″')
    words = pdf_words(pdf, T) if info['pages'] else None
    imgs = pdf_images(pdf, T) if info['pages'] else None
    lay, note = html_layout(src, theme, work, layout_file)
    if note:
        run['limits'].append(note)
    png = None
    if T['pdftoppm'] and info['pages']:
        png = work / 'png'
        shutil.rmtree(png, ignore_errors=True)
        png.mkdir(parents=True)
        sh([T['pdftoppm'], '-png', '-scale-to', '1920', str(pdf), str(png / 'p')], timeout=180)
        run['png'] = str(png.relative_to(src))
    pdf_font_names, fams = set(), set()
    for s in slides:
        fams.update(s['fonts'])
        row = {'n': s['n'], 'hidden': s['hidden'], 'issues': [], 'skipped': []}
        pg = pages.get(s['n'])
        if not pg:
            row['status'] = 'not_run' if s['hidden'] or not info['pages'] else 'partial'
            row['skipped'].append('скрытый слайд не экспортирован' if s['hidden'] else 'страница PDF не найдена')
            run['slides'].append(row)
            continue
        fonts = pdf_fonts(pdf, T, pg)
        pdf_font_names.update(f['pdf_name'] for f in fonts or [])
        iss, skipped = check_page(s, (words or {}).get(pg, []) if words is not None else None, dict(info, n=pg), fonts, imgs, lay.get(s['n'], []) if lay is not None else None)
        row.update(page=pg, issues=[{'level': l, 'text': t} for l, t in iss], skipped=skipped, fonts=fonts or [])
        if png:
            cand = sorted(png.glob(f'p-*{pg}.png'), key=lambda p: len(p.name))
            cand = [c for c in cand if int(re.search(r'-(\d+)\.png$', c.name).group(1)) == pg]
            if cand:
                dst = png / f's{s["n"]}.png'
                cand[0].rename(dst)
                ref = png_of(src, theme, s['n'])
                if ref.is_file():
                    row['diff_html'] = visual_diff(dst, ref)
        row['status'] = 'failed' if any(l == 'err' for l, _ in iss) else 'partial' if skipped else 'passed'
        run['slides'].append(row)
    if run['limits']:          # проверка неполная (нет poppler, Chrome, шрифта в замере): «пройдено» только частично
        for row in run['slides']:
            if row['status'] == 'passed':
                row['status'] = 'partial'
    run['fonts'] = installed_versions(sorted(fams), sorted(pdf_font_names), font_index(app.get('fonts', ())))
    run['pdf_fonts'] = sorted(pdf_font_names)
    sts = [r['status'] for r in run['slides'] if not r['hidden']]
    run['status'] = ('failed' if run['errors'] or 'failed' in sts else
                     'partial' if (not sts or 'partial' in sts or 'not_run' in sts or run['limits']) else 'passed')
    return run


# ───────── итог: прогоны + ручные записи, свежесть по sha ─────────
def load_state(out):
    p = out / 'target_check.json'
    st = json.load(open(p, encoding='utf-8')) if p.is_file() else {'version': 1, 'runs': {}, 'manual': []}
    leg = out / 'powerpoint.json'      # запись просмотра до 1.8: out/powerpoint.json
    if leg.is_file() and not st.get('legacy_imported'):
        pp = json.load(open(leg, encoding='utf-8'))
        for v in pp.get('viewed', []):
            st['manual'].append({'theme': v.get('theme', 'dark'), 'slides': v.get('slides', []), 'bad': [], 'by': v.get('by', '?'),
                                 'app': 'Microsoft PowerPoint', 'notes': pp.get('notes', ''), 'when': None, 'pptx_sha': None, 'legacy': True})
    return st


def summarize(src, st):
    name, themes = project(src)
    shas = {th: sha(p) for th, p in themes.items()}
    count = {}
    for th, p in themes.items():
        try:
            m = pptx_model(p)
            count[th] = [s['n'] for s in m['slides'] if not s['hidden']]
        except Exception:
            count[th] = []
    S = {'status': 'not_run', 'themes': {}, 'slides': {}, 'errors': [], 'warnings': [], 'items': []}
    stat = []
    for th, p in themes.items():
        T = {'pptx': p.name, 'auto': None, 'manual': [], 'status': 'not_run'}
        r = st.get('runs', {}).get(th)
        if r:
            stale = r.get('pptx_sha') != shas[th]
            T['auto'] = {k: r.get(k) for k in ('status', 'app', 'app_version', 'when', 'reason', 'pdf', 'png', 'limits', 'pdf_fonts', 'fonts')}
            T['auto']['stale'] = stale
            if stale:
                S['warnings'].append(f'{THEMES[th]}: проверка в {r.get("app")} от {r.get("when")} относится к прошлой сборке PPTX — повторите run')
            elif r.get('status') != 'not_run':
                for row in r.get('slides', []):
                    d = S['slides'].setdefault(str(row['n']), {}).setdefault(th, {})
                    d['auto'] = row['status']
                    d['app'] = r.get('app')
                    for i in row.get('issues', []):
                        (S['errors'] if i['level'] == 'err' else S['warnings']).append(f'{THEMES[th]}, слайд {row["n"]}: {i["text"]}')
                S['errors'] += [f'{THEMES[th]}: {e}' for e in r.get('errors', [])]
        seen, bad = set(), set()
        for m in st.get('manual', []):
            if m.get('theme') != th:
                continue
            stale = m.get('pptx_sha') is not None and m.get('pptx_sha') != shas[th]
            T['manual'].append({k: m.get(k) for k in ('slides', 'bad', 'by', 'app', 'when', 'notes', 'legacy')} | {'stale': stale})
            if stale:
                S['warnings'].append(f'{THEMES[th]}: ручной просмотр слайдов {m.get("slides")} от {m.get("when")} — до пересборки PPTX')
                continue
            seen.update(m.get('slides', []))
            bad.update(m.get('bad', []))
        for n in sorted(seen):
            d = S['slides'].setdefault(str(n), {}).setdefault(th, {})
            d['manual'] = 'failed' if n in bad else 'passed'
        S['errors'] += [f'{THEMES[th]}, слайд {n}: проблема при ручном просмотре' for n in sorted(bad)]
        # статус темы: слайд подтверждён прогоном или ручным просмотром; ошибка где угодно — ошибки
        auto = T['auto']['status'] if T['auto'] and not T['auto']['stale'] else None
        per = []
        for n in count[th]:
            d = S['slides'].get(str(n), {}).get(th, {})
            v = [x for x in (d.get('auto'), d.get('manual')) if x and x != 'not_run']
            per.append('failed' if 'failed' in v else 'passed' if 'passed' in v else 'partial' if v else 'not_run')
        if 'failed' in per or auto == 'failed':
            T['status'] = 'failed'
        elif per and all(x == 'passed' for x in per):
            T['status'] = 'passed'
        elif any(x != 'not_run' for x in per):
            T['status'] = 'partial'
        T['slides_total'] = len(count[th])
        T['slides_ok'] = per.count('passed')
        S['themes'][th] = T
        stat.append(T['status'])
    if not themes:
        S['errors'].append('PPTX рядом с проектом не найден: соберите деку (node build.js)')
        S['status'] = 'not_run'
    elif 'failed' in stat:
        S['status'] = 'failed'
    elif all(x == 'passed' for x in stat):
        S['status'] = 'passed'
    elif any(x != 'not_run' for x in stat):
        S['status'] = 'partial'
    for th, T in S['themes'].items():
        a = T['auto']
        how = []
        if a and not a['stale']:
            who = f"{a['app']} {a.get('app_version') or ''}".strip() if a['app'] not in (None, '—') else 'автопроверка'
            how.append(f"{who}: {LABEL[a['status']]}" + (f" ({a['reason']})" if a.get('reason') else ''))
        mv = sorted({n for m in T['manual'] if not m['stale'] for n in m.get('slides') or []})
        if mv:
            how.append(f"вручную слайды {', '.join(map(str, mv))}")
        S['items'].append(f"{THEMES[th]} ({T['pptx']}): {LABEL[T['status']]}; слайдов подтверждено {T['slides_ok']} из {T['slides_total']}"
                          + (f" — {'; '.join(how)}" if how else ''))
    S['label'] = LABEL[S['status']]
    return S


def save(out, st, summary):
    st['summary'] = summary
    st['legacy_imported'] = st.get('legacy_imported') or (out / 'powerpoint.json').is_file()
    json.dump(st, open(out / 'target_check.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)


def report(S):
    for e in S['errors']:
        print('[ERR]', e)
    for w in S['warnings']:
        print('[WARN]', w)
    for i in S['items']:
        print('  ' + i)
    print(f"PPTX в целевом приложении: {S['label']}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('cmd', choices=['run', 'manual', 'status', 'env'])
    ap.add_argument('src', nargs='?')
    ap.add_argument('--app', default='auto', help='run: auto|powerpoint|libreoffice|none; manual: чем смотрели')
    ap.add_argument('--theme', default='all')
    ap.add_argument('--slides', help='manual: номера просмотренных слайдов, 1,3,5 или 1-12')
    ap.add_argument('--bad', default='', help='manual: слайды, где при просмотре нашлась проблема')
    ap.add_argument('--by', default='человек')
    ap.add_argument('--notes', default='')
    ap.add_argument('--timeout', type=int, default=240)
    ap.add_argument('--json', help='status: итог файлом')
    ap.add_argument('--fake-pdf', help='тесты: готовый PDF вместо приложения (адаптер fake)')
    ap.add_argument('--html-layout', help='тесты: замер HTML-двойника файлом вместо Chrome')
    a = ap.parse_args()
    if a.cmd == 'env':
        E = environment()
        for ad in E['adapters']:
            v = ad.get('version') or (lo_version(ad) if ad['name'] == 'libreoffice' and ad['available'] else '')
            print(f"  {ad['title']}: " + (f'есть {v}'.strip() if ad['available'] else ad['why']))
        print('  poppler: ' + ', '.join(f'{t} {"есть" if p else "нет"}' for t, p in E['tools'].items()))
        print('  Chrome для замера HTML-двойника: ' + ('есть' if E['chrome'] else 'нет'))
        return 0
    if not a.src:
        sys.exit('укажите папку проекта src-<тема>')
    src = Path(a.src).resolve()
    out = src / 'out'
    if not out.is_dir():
        sys.exit(f'нет {out}: соберите деку (node build.js)')
    name, themes = project(src)
    st = load_state(out)
    want = list(themes) if a.theme == 'all' else [a.theme]
    for th in want:
        if th not in THEMES:
            sys.exit(f'--theme: dark, light или all, а не {th}')
        if th not in themes:
            sys.exit(f'нет PPTX темы «{THEMES[th]}» рядом с проектом: соберите её ({"THEME=light " if th == "light" else ""}node build.js)')

    if a.cmd == 'manual':
        if a.theme == 'all' or not a.slides:
            sys.exit('manual: укажите одну тему и слайды: --theme dark --slides 1,3,5')
        def nums(s):
            r = set()
            for part in filter(None, s.replace(' ', '').split(',')):
                lo, _, hi = part.partition('-')
                r.update(range(int(lo), int(hi or lo) + 1))
            return sorted(r)
        total = len(pptx_model(themes[a.theme])['slides'])
        sl, bad = nums(a.slides), nums(a.bad)
        if not sl or max(sl) > total or min(sl) < 1:
            sys.exit(f'--slides: номера от 1 до {total}')
        if not set(bad) <= set(sl):
            sys.exit('--bad: только среди просмотренных слайдов')
        app = {'auto': 'Microsoft PowerPoint', 'powerpoint': 'Microsoft PowerPoint', 'libreoffice': 'LibreOffice', 'keynote': 'Keynote'}.get(a.app, a.app)
        st['manual'].append({'theme': a.theme, 'slides': sl, 'bad': bad, 'by': a.by, 'app': app, 'notes': a.notes,
                             'when': now(), 'pptx_sha': sha(themes[a.theme])})
        S = summarize(src, st)
        save(out, st, S)
        report(S)
        return 1 if S['status'] == 'failed' else 0

    if a.cmd == 'run':
        ads = {'powerpoint': powerpoint(), 'libreoffice': libreoffice()}
        if a.app == 'fake':
            if not a.fake_pdf:
                sys.exit('--app fake только для тестов и с --fake-pdf')
            ad = {'name': 'fake', 'title': 'тестовый адаптер', 'available': True, 'version': '0', 'pdf': a.fake_pdf}
        elif a.app == 'auto':
            ad = next((x for x in ads.values() if x['available']), None)   # не ответил — следующий доступный
        elif a.app in ads:
            ad = ads[a.app]
        elif a.app == 'none':
            ad = None
        else:
            sys.exit('--app: auto, powerpoint, libreoffice или none')
        for th in want:
            pptx = themes[th]
            work = out / 'target' / th
            work.mkdir(parents=True, exist_ok=True)
            base = {'theme': th, 'pptx': pptx.name, 'pptx_sha': sha(pptx), 'when': now(), 'slides': [], 'errors': [], 'limits': []}
            if ad is None or not ad['available']:
                why = ad['why'] if ad else ('проверка отключена (--app none)' if a.app == 'none' else
                                            'нет ни PowerPoint, ни LibreOffice: ' + '; '.join(x['why'] for x in ads.values()))
                st['runs'][th] = base | {'status': 'not_run', 'app': ad['title'] if ad else '—', 'adapter': ad['name'] if ad else None, 'reason': why}
                print(f'  {THEMES[th]}: не выполнялось — {why}')
                continue
            chain = [x for x in ads.values() if x['available']] if a.app == 'auto' else [ad]
            pdf = work / f'{pptx.stem}.pdf'
            fails_ = []
            for ad in chain:
                if ad['name'] == 'libreoffice' and not ad.get('version'):
                    ad['version'] = lo_version(ad)
                pdf.unlink(missing_ok=True)
                print(f'  {THEMES[th]}: {ad["title"]} {ad.get("version") or ""} → PDF…')
                err = {'powerpoint': export_powerpoint, 'libreoffice': export_libreoffice, 'fake': export_fake}[ad['name']](ad, pptx, pdf, a.timeout)
                if not err:
                    break
                fails_.append(f'{ad["title"]}: {err}')
                print(f'  {THEMES[th]}: {ad["title"]} — {err}')
            if err:
                st['runs'][th] = base | {'status': 'not_run', 'app': ad['title'], 'adapter': ad['name'], 'app_version': ad.get('version'), 'reason': '; '.join(fails_)}
                print(f'  {THEMES[th]}: не выполнялось')
                continue
            r = analyze(src, th, pptx, pdf, ad, work, a.html_layout)
            st['runs'][th] = r
            bad = sum(1 for x in r['slides'] if x['status'] == 'failed')
            print(f'  {THEMES[th]}: {LABEL[r["status"]]} — слайдов {len(r["slides"])}, с ошибками {bad}; PDF {r["pdf"]}')
            for lim in r['limits']:
                print(f'    не проверено: {lim}')
        S = summarize(src, st)
        save(out, st, S)
        report(S)
        return 1 if S['status'] == 'failed' else 0

    S = summarize(src, st)
    if (out / 'target_check.json').is_file():
        save(out, st, S)
    if not st.get('runs') and not st.get('manual'):
        av = [x['title'] for x in (powerpoint(), libreoffice()) if x['available']]
        S['items'].append('не запускалась: ' + (f'есть {", ".join(av)} — pptx_target_check.sh run <src>' if av else
                                                'нет PowerPoint и LibreOffice; ручной просмотр — pptx_target_check.sh manual <src> --theme dark --slides 1-12'))
    if a.json:
        json.dump(S, open(a.json, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    report(S)
    return 1 if S['status'] == 'failed' else 0


if __name__ == '__main__':
    sys.exit(main())

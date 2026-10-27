# Аудит готового PPTX: сетка, поля, выравнивание, типографика, цвет, признаки ИИ-слопа в тексте.
# Запуск: $SKILL/.venv/bin/python -I pptx_audit.py <deck.pptx> [--theme theme.json] [--json report.json] [--strict]
# Код выхода 1, если есть ошибки (ERR). Предупреждения (WARN) не валят проверку, кроме --strict.
import argparse
import colorsys
import io
import json
import os
import re
import sys
from collections import defaultdict

from pptx import Presentation
from pptx.util import Emu

EMU = 914400
ap = argparse.ArgumentParser()
ap.add_argument('pptx')
ap.add_argument('--theme', default=None, help='theme.json проекта: шрифты, поля, минимальный кегль')
ap.add_argument('--json', default=None)
ap.add_argument('--strict', action='store_true')
ap.add_argument('--manifest', default=None, help='out/slides.json проекта: архетипы слайдов (по умолчанию рядом с theme.json)')
a = ap.parse_args()

def find_theme():
    if a.theme:
        return a.theme
    # build.js кладёт pptx рядом с папкой проекта: ищем src-*/deck.json, чьё name совпадает с именем файла
    here = os.path.dirname(os.path.abspath(a.pptx))
    stem = os.path.splitext(os.path.basename(a.pptx))[0]
    for d in sorted(os.listdir(here)):
        dj = os.path.join(here, d, 'deck.json')
        if os.path.isfile(dj) and stem.startswith(json.load(open(dj)).get('name', '\0')):
            return os.path.join(here, d, 'theme.json')
    return None


tp = find_theme()
theme = json.load(open(tp)) if tp and os.path.isfile(tp) else {}
print('тема:', tp or 'не найдена, проверки шрифтов и полей по умолчанию')
FONTS = {v for k, v in theme.get('font', {}).items() if isinstance(v, str)} or None
MARGIN = theme.get('canvas', {}).get('margin', 0.5)
MIN_PT = theme.get('type', {}).get('min', 8)
# Архетипы из манифеста сборки: текстовые раскладки (archetype «text-…») освобождены от бана простых слайдов,
# вместо него проверяется объём текста
mp = a.manifest or (os.path.join(os.path.dirname(tp), 'out', 'slides.json') if tp else None)
ARCH = {e['n']: e.get('archetype', '') for e in json.load(open(mp))} if mp and os.path.isfile(mp) else {}
TEXT_WORDS_MAX = 90

prs = Presentation(a.pptx)
SW, SH = prs.slide_width / EMU, prs.slide_height / EMU
issues = []
stats = []      # по слайдам для --json: картинки, графика, контейнеры, архетип (eval_deck.py: визуальная плотность)


def add(level, slide, msg):
    issues.append({'level': level, 'slide': slide, 'msg': msg})


SLOP_WORDS = r'бесшовн|инновацион|уникальн|революцион|синерги|передов|непревзойд|беспрецедент|в рамках|осуществлен|позволяет обеспечить|является ключев|играет ключев|ключевую роль'
GENERIC_STEPS = r'^(шаг|этап|фаза|stage|step|phase)\s*0?\d+\s*$'


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6 and hasattr(sh, 'shapes'):   # группа
            yield from walk(sh.shapes)
        else:
            yield sh


def color_of(run):
    try:
        if run.font.color and run.font.color.type is not None and run.font.color.rgb is not None:
            return str(run.font.color.rgb)
    except AttributeError:
        pass
    return None


def saturated(hexs):
    r, g, b = (int(hexs[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    return s > 0.35 and 0.15 < l < 0.9, round(h * 12) % 12


for n, slide in enumerate(prs.slides, 1):
    if slide._element.get('show') == '0':
        add('INFO', n, 'скрытый слайд')
    sizes, accents, lefts, texts = set(), set(), [], []
    covered, content_bottom, content_right = 0.0, 0.0, 0.0
    body = []   # фигуры тела слайда для проверки «простого слайда»: (вид, x, y, w, h, текст)
    BODY_TOP, BODY_BOT = theme.get('canvas', {}).get('bodyTop', 1.9), theme.get('canvas', {}).get('bodyBottom', SH - 0.7)
    for sh in walk(slide.shapes):
        if sh.left is None or sh.width is None:
            continue
        x, y, w, h = sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU
        full_bleed = w >= SW * 0.95 and h >= SH * 0.95
        if not full_bleed and (x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01):
            add('ERR', n, f'«{sh.name}» выходит за холст: x {x:.2f} y {y:.2f} w {w:.2f} h {h:.2f}')
        if not full_bleed and y > BODY_TOP - 0.05 and y < BODY_BOT and w * h < SW * SH * 0.4:
            txt = sh.text_frame.text.strip() if sh.has_text_frame else ''
            kind = 'image' if sh.shape_type == 13 else 'chart' if getattr(sh, 'has_chart', False) else 'text' if txt else 'line' if (w < 0.02 or h < 0.02) else 'dot' if (w < 0.13 and h < 0.13) else 'shape'
            body.append((kind, x, y, w, h, txt))
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip()
            if not full_bleed and (x < MARGIN - 0.06 or x + w > SW - MARGIN + 0.06) and w < SW * 0.9:
                add('WARN', n, f'текст «{t[:30]}» заходит в поле {MARGIN}″ (x {x:.2f}, правый край {x + w:.2f})')
            # линия выравнивания зависит от выключки: левый край, центр или правый край
            al = next((p.alignment for p in sh.text_frame.paragraphs if p.alignment is not None), None)
            anchor = {2: x + w / 2, 3: x + w}.get(int(al) if al is not None else 1, x)
            lefts.append((round(anchor, 3), t[:24], {2: 'центр', 3: 'правый край'}.get(int(al) if al is not None else 1, 'левый край')))
            texts.append((x, y, w, h, t))
            for p in sh.text_frame.paragraphs:
                pt = p.text.strip()
                if not pt:
                    continue
                if re.match(r'^[—–-]\s', pt):
                    add('WARN', n, f'тире вместо маркера: «{pt[:40]}»')
                if pt.count('·') > 1:
                    add('WARN', n, f'больше одной средней точки в строке: «{pt[:50]}»')
                if re.match(r'^0\d\b', pt) and len(pt) <= 4:
                    add('WARN', n, f'декоративная нумерация «{pt}»')
                if re.match(GENERIC_STEPS, pt, re.I):
                    add('WARN', n, f'безликая подпись шага «{pt}»: назовите шаг по сути')
                m = re.search(SLOP_WORDS, pt, re.I)
                if m:
                    add('WARN', n, f'слово-штамп «{m.group(0)}»: «{pt[:60]}»')
                if pt.startswith('[экран'):
                    add('WARN', n, f'мокап экрана не вставлен: «{pt[:40]}» — положите PNG экрана в img/ и передайте путь в data')
                if '  ' in pt:
                    add('WARN', n, f'двойной пробел: «{pt[:40]}»')
                for r in p.runs:
                    if r.font.size:
                        s = r.font.size.pt
                        sizes.add(s)
                        if s < MIN_PT:
                            add('ERR', n, f'кегль {s} pt меньше {MIN_PT}: «{r.text[:30]}»')
                    if FONTS and r.font.name and r.font.name not in FONTS:
                        add('ERR', n, f'шрифт «{r.font.name}» не из темы {sorted(FONTS)}: «{r.text[:30]}»')
                    c = color_of(r)
                    if c:
                        sat, hue = saturated(c)
                        if sat:
                            accents.add(hue)
        try:
            if sh.fill.type == 1:
                sat, hue = saturated(str(sh.fill.fore_color.rgb))
                if sat:
                    accents.add(hue)
        except (AttributeError, TypeError, ValueError, KeyError):
            pass
        if not full_bleed and w * h < SW * SH * 0.4:
            covered += w * h
            if y > BODY_TOP - 0.05 and y < BODY_BOT:                       # содержимое тела, без шапки и подвала
                content_bottom = max(content_bottom, y + (h if not (sh.has_text_frame and sh.text_frame.text.strip()) else min(h, 0.6)))
                if sh.has_text_frame and sh.text_frame.text.strip():
                    longest = max((len(l) for l in sh.text_frame.text.split('\n')), default=0)
                    pts = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size] or [12]
                    content_right = max(content_right, x + min(w, longest * max(pts) * 0.62 / 72))
                else:
                    content_right = max(content_right, x + w)

    if len(sizes) > 8:
        add('WARN', n, f'{len(sizes)} разных кеглей на слайде ({sorted(sizes)}): шапка уже даёт 5, на тело оставьте 2–3')
    if len(accents) > 2:
        add('WARN', n, f'{len(accents)} насыщенных оттенков на слайде: акцент должен быть один (плюс семантический, если нужен)')
    # почти совпадающие левые края: выглядит как ошибка выравнивания
    for kind in {k for _, _, k in lefts}:
        xs = sorted({x for x, _, k in lefts if k == kind})
        for i in range(1, len(xs)):
            d = xs[i] - xs[i - 1]
            if 0.008 < d < 0.045:
                who = [t for x, t, k in lefts if k == kind and x in (xs[i - 1], xs[i])][:2]
                add('WARN', n, f'линии «{kind}» {xs[i - 1]:.3f}″ и {xs[i]:.3f}″ почти совпадают: выровняйте ({who})')
    # пересечение текстовых боксов
    for i in range(len(texts)):
        for j in range(i + 1, len(texts)):
            x1, y1, w1, h1, t1 = texts[i]
            x2, y2, w2, h2, t2 = texts[j]
            ox = min(x1 + w1, x2 + w2) - max(x1, x2)
            oy = min(y1 + h1, y2 + h2) - max(y1, y2)
            if ox > 0.05 and oy > 0.05 and ox * oy > 0.02:
                add('WARN', n, f'тексты пересекаются: «{t1[:20]}» и «{t2[:20]}»')
    # Бан простых слайдов (reference/08-infographics.md): графика = фигуры без текста и фигуры-ячейки с коротким числом
    texts_b = [b for b in body if b[0] == 'text']
    def holds(b):
        return [t for t in texts_b if t[1] >= b[1] - 0.01 and t[2] >= b[2] - 0.01 and t[1] + t[3] <= b[1] + b[3] + 0.01 and t[2] + t[4] <= b[2] + b[4] + 0.01]
    # картинки (3D, экраны) считаются по всему слайду: у обложки нет шапки, объект может начинаться выше тела
    images_n = sum(1 for sh in walk(slide.shapes) if sh.shape_type == 13 and sh.width is not None
                   and not (sh.width / EMU >= SW * 0.95 and sh.height / EMU >= SH * 0.95))
    for sh in walk(slide.shapes):
        if sh.shape_type == 13 and sh.width is not None and sh.width / EMU < SW * 0.95:
            content_right = max(content_right, (sh.left + sh.width) / EMU)
    graphics, containers = 0, 0
    conts = []             # карточки и панели с текстом: (x, y, w, h)
    for b in body:
        if b[0] == 'chart':
            graphics += 2          # нативная диаграмма PowerPoint — инфографика сама по себе
            continue
        if b[0] != 'shape':
            continue
        inner = holds(b)
        if not inner or (len(inner) <= 2 and all(len(t[5]) <= 14 for t in inner)):
            graphics += 1          # полоса, сегмент, отметка, ячейка матрицы с числом
        else:
            containers += 1        # карточка или панель с текстом
            conts.append(b[1:5])
    # линии не считаются: разделители и оси не кодируют данные сами по себе
    if ARCH.get(n, '').startswith('text-'):
        words = sum(len(t[5].split()) for t in body if t[0] == 'text')
        if words > TEXT_WORDS_MAX:
            add('WARN', n, f'стена текста: {words} слов в теле текстового слайда (не больше {TEXT_WORDS_MAX}): сократите или разнесите на два слайда')
    elif body and not slide._element.get('show') == '0':
        if images_n == 0 and graphics < 2 and containers == 0:
            add('ERR', n, 'простой слайд: только текст, нет инфографики, 3D-объекта или экрана (бан простых слайдов)')
        # Запрет блоковых слайдов (правило пользователя 27.10.2026, reference/10-beyond-blocks.md): ряд или сетка одинаковых
        # карточек с текстом без связей между ними — типичный блоковый слайд; 4+ контейнера — блочная раскладка. Схемы связей
        # библиотеки (архетип scheme-*) — исключение: их узлы соединены линиями.
        same = []
        for c in conts:
            grp = [d for d in conts if abs(d[2] - c[2]) <= 0.08 * c[2] and abs(d[3] - c[3]) <= 0.08 * c[3]]
            if len(grp) > len(same):
                same = grp
        aligned = len(same) >= 3 and (len({round(d[1], 1) for d in same}) < len(same) or len({round(d[0], 1) for d in same}) < len(same))
        if ARCH.get(n, '').startswith('scheme-'):
            graphics += 2          # схема связей (reference/19-schemes.md): узлы-карточки соединены линиями, это не блочная раскладка
        elif aligned:
            add('ERR', n, f'блоковый слайд: {len(same)} одинаковых карточки с текстом рядом или сеткой. Запрещено: замените схемой '
                          f'(viz_suggest.py deck), инфографикой или раскладкой без контейнеров (reference/10-beyond-blocks.md)')
        elif containers > 3:
            add('ERR', n, f'блочная раскладка: {containers} контейнеров с текстом. Замените часть сеткой, линиями и воздухом (reference/10-beyond-blocks.md)')
        if images_n == 0 and graphics < 2 and containers == 0:
            pass
        elif images_n == 0 and graphics < 2:
            add('WARN', n, f'бедный слайд: {containers} контейнер(а) с текстом, нет инфографики и 3D-объекта. Добавьте кодирование данных или контекстный объект')
    stats.append({'n': n, 'hidden': slide._element.get('show') == '0', 'archetype': ARCH.get(n, ''), 'images': images_n,
                  'graphics': graphics, 'containers': containers})
    body_h = BODY_BOT - BODY_TOP
    if content_bottom and (BODY_BOT - content_bottom) / body_h > 0.3:
        add('WARN', n, f'нижние {(BODY_BOT - content_bottom) / body_h:.0%} тела пустые (содержимое кончается на {content_bottom:.2f}″): растяните сетку или перекомпонуйте')
    body_w = SW - 2 * MARGIN
    if content_right and (SW - MARGIN - content_right) / body_w > 0.35:
        add('WARN', n, f'правые {(SW - MARGIN - content_right) / body_w:.0%} тела пустые (содержимое кончается на {content_right:.2f}″)')
    fill = covered / (SW * SH)
    if fill < 0.12 and len(texts) > 1:
        add('WARN', n, f'занято ~{fill:.0%} холста: проверьте пустоту без функции')

# Нативные диаграммы PowerPoint (kit/charts.js → nativeChart): встроенная таблица без внешней связи, ссылки ряда
# совпадают с кэшем диаграммы (иначе «Изменить данные» покажет другие числа), шрифты и кегли из темы.
try:
    import openpyxl as _xl
except ImportError:
    _xl = None


def ref_values(ws, f):
    # 'Sheet1!$B$2:$B$7' → значения ячеек по порядку
    m = re.match(r"^'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", f or '')
    if not m:
        return None
    c1, r1, c2, r2 = m.group(2), int(m.group(3)), m.group(4) or m.group(2), int(m.group(5) or m.group(3))
    return [c.value for row in ws[f'{c1}{r1}:{c2}{r2}'] for c in row]


def same(a, b):
    if isinstance(a, (int, float)) or isinstance(b, (int, float)):
        try:
            return abs(float(a) - float(b)) < 1e-9
        except (TypeError, ValueError):
            return False
    return (a if a is not None else '') == (b if b is not None else '')


for n, slide in enumerate(prs.slides, 1):
    for sh in walk(slide.shapes):
        if not getattr(sh, 'has_chart', False):
            continue
        cs, name = sh.chart._chartSpace, f'диаграмма «{sh.name}»'
        for tf in sorted({t for t in cs.xpath('.//a:latin/@typeface') if not t.startswith('+')}):
            if FONTS and tf not in FONTS:
                add('ERR', n, f'{name}: шрифт «{tf}» не из темы {sorted(FONTS)}')
        for sz in {int(v) for v in cs.xpath('.//a:defRPr/@sz')}:
            if sz / 100 < MIN_PT:
                add('ERR', n, f'{name}: кегль {sz / 100} pt меньше {MIN_PT}')
        ext = [r for r in sh.chart.part.rels.values() if r.is_external]
        if ext:
            add('ERR', n, f'{name}: связь с внешним файлом ({ext[0].target_ref}): данные должны быть встроены')
        wb = sh.chart.part.chart_workbook.xlsx_part
        if wb is None:
            add('ERR', n, f'{name}: нет встроенной таблицы, «Изменить данные» в PowerPoint не откроется')
            continue
        if _xl is None:
            continue
        try:
            book = _xl.load_workbook(io.BytesIO(wb.blob))
        except Exception as e:
            add('ERR', n, f'{name}: встроенная таблица не читается ({e.__class__.__name__})')
            continue
        bad = []
        for ser in cs.xpath('.//c:ser'):
            for part, cache in (('c:tx', './/c:strCache//c:v'), ('c:cat', './/c:lvl/c:pt/c:v|.//c:strCache//c:v|.//c:numCache//c:v'), ('c:val', './/c:numCache/c:pt/c:v')):
                fs = ser.xpath(f'./{part}//c:f')
                if not fs:
                    continue
                f = fs[0].text
                sheet = f.split('!')[0].strip("'")
                if sheet not in book.sheetnames:
                    bad.append(f'{f}: нет листа')
                    continue
                got = ref_values(book[sheet], f)
                want = [v.text for v in ser.xpath(f'./{part}')[0].xpath(cache)]
                if part == 'c:val':
                    pts = {int(p.get('idx')): p.xpath('./c:v')[0].text for p in ser.xpath('./c:val//c:numCache/c:pt')}
                    want = [pts.get(i) for i in range(len(got or []))]
                if got is None or len(got) != len(want) or not all(same(g, w) for g, w in zip(got, want)):
                    bad.append(f'{f}: таблица {got} ≠ кэш {want}')
        if bad:
            add('ERR', n, f'{name}: кэш диаграммы расходится со встроенной таблицей: {bad[0]}')

# Размещение 3D-объектов и картинок (reference/09-3d-objects.md, раздел 6)
try:
    from PIL import Image as _PIL
except ImportError:
    _PIL = None
stripes = theme.get('bg', {}).get('stripes')
STRIPE_X = min(st[0] for st in stripes) * SW if stripes else None
series = []   # (slide, cx, cy, видимая площадь)


def ink_box(sh):
    # фактическая область текста: ширина по самой длинной строке, высота по числу строк
    x, y, w, h = sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU
    pts = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size] or [12]
    lines = [l for l in sh.text_frame.text.split('\n')] or ['']
    iw = min(w, max(len(l) for l in lines) * max(pts) * 0.62 / 72)
    ih = min(h, len(lines) * max(pts) * 1.25 / 72)
    al = next((p.alignment for p in sh.text_frame.paragraphs if p.alignment is not None), None)
    ix = x + (w - iw) / 2 if al is not None and int(al) == 2 else x + w - iw if al is not None and int(al) == 3 else x
    va = sh.text_frame.vertical_anchor
    iy = y + (h - ih) / 2 if va is not None and int(va) == 3 else y + h - ih if va is not None and int(va) == 4 else y
    return ix, iy, iw, ih


def gap(a, b):
    dx = max(b[0] - (a[0] + a[2]), a[0] - (b[0] + b[2]), 0)
    dy = max(b[1] - (a[1] + a[3]), a[1] - (b[1] + b[3]), 0)
    return (dx * dx + dy * dy) ** 0.5


for n, slide in enumerate(prs.slides, 1):
    shapes = [sh for sh in walk(slide.shapes) if sh.width is not None]
    pics = [sh for sh in shapes if sh.shape_type == 13 and sh.width / EMU < SW * 0.95]
    if not pics:
        continue
    inks = [ink_box(sh) for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    anchors = inks + [(sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU) for sh in shapes
                      if sh.shape_type != 13 and not (sh.has_text_frame and sh.text_frame.text.strip())
                      and sh.width / EMU * sh.height / EMU < SW * SH * 0.3 and sh.width / EMU > 0.15 and sh.height / EMU > 0.15]
    cover = n == 1 or not any(i[1] > 1.5 for i in inks)
    for pic in pics:
        b = (pic.left / EMU, pic.top / EMU, pic.width / EMU, pic.height / EMU)
        for t in inks:
            ox = min(b[0] + b[2], t[0] + t[2]) - max(b[0], t[0])
            oy = min(b[1] + b[3], t[1] + t[3]) - max(b[1], t[1])
            if ox > 0.05 and oy > 0.05 and ox * oy > 0.05:
                add('ERR', n, f'объект перекрывает текст на {ox:.2f}×{oy:.2f}″')
                break
        contains = lambda c, r: c[0] <= r[0] + 0.01 and c[1] <= r[1] + 0.01 and c[0] + c[2] >= r[0] + r[2] - 0.01 and c[1] + c[3] >= r[1] + r[3] - 0.01
        inside = [c for c in anchors if contains(c, b)]
        cands = [c for c in anchors if c not in inside]
        nearest = min(cands, key=lambda c: gap(b, c), default=None)
        near = gap(b, nearest) if nearest else 99
        # граница панели между объектом и ближайшим содержимым — тоже отрыв: объект снаружи, содержимое внутри
        panels = [(sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU) for sh in shapes
                  if sh.shape_type != 13 and sh.width / EMU * sh.height / EMU >= SW * SH * 0.3 and sh.width / EMU < SW * 0.95]
        across = nearest is not None and any(contains(pn, nearest) and not contains(pn, b) for pn in panels)
        if not cover and not inside and (near > 0.5 or across):
            add('WARN', n, f'объект оторван от содержания: до ближайшего элемента {near:.2f}″. Поставьте его у якоря (внутри тайла, над подписью, у числа)')
        on_panel = any(contains(pn, b) for pn in panels)   # панель закрывает полосы: объект на ней не спорит с фоном
        if STRIPE_X and not cover and not on_panel and b[0] + b[2] > STRIPE_X + 0.2 * b[2] and b[1] < SH * 0.6:   # на обложке объект на фоне полос — осознанная композиция
            add('WARN', n, f'объект заходит в зону фирменных полос (x > {STRIPE_X:.2f}″)')
        if b[0] + b[2] > SW - MARGIN - 0.7 and b[1] < 1.2:
            add('WARN', n, 'объект в зоне знака бренда')
        fill = 1.0
        if _PIL is not None:
            try:
                im = _PIL.open(io.BytesIO(pic.image.blob))
                if im.mode in ('RGBA', 'LA'):
                    al = im.getchannel('A')
                    fill = sum(al.histogram()[129:]) / (im.width * im.height)   # доля непрозрачных пикселей (alpha > 128)
            except Exception:
                pass
        series.append((n, b[0] + b[2] / 2, b[1] + b[3] / 2, b[2] * b[3] * fill))

# линии (выноски, оси, разделители), которые перечёркивают текст
for n, slide in enumerate(prs.slides, 1):
    shapes = [sh for sh in walk(slide.shapes) if sh.width is not None]
    inks = [(ink_box(sh), sh.text_frame.text.strip()[:30]) for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    def knocked(li, box):
        # непрозрачная фигура выше линии по слоям закрывает её под текстом (подложка числа)
        for sh2 in shapes[li + 1:]:
            try:
                solid = sh2.fill.type == 1 and not (sh2.has_text_frame and sh2.text_frame.text.strip())
            except (AttributeError, TypeError):
                solid = False
            if solid:
                bx, by, bw, bh = sh2.left / EMU, sh2.top / EMU, sh2.width / EMU, sh2.height / EMU
                if bx <= box[0] + 0.02 and by <= box[1] + 0.02 and bx + bw >= box[0] + box[2] - 0.02 and by + bh >= box[1] + box[3] - 0.02:
                    return True
        return False
    for li, sh in enumerate(shapes):
        if sh.has_text_frame and sh.text_frame.text.strip():
            continue
        x, y, w, h = sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU
        if not ((w < 0.03 and h > 0.3) or (h < 0.03 and w > 0.3)):
            continue
        for (ix, iy, iw, ih), t in inks:
            if knocked(li, (ix, iy, iw, ih)):
                continue
            if w < 0.03 and ix + 0.03 < x < ix + iw - 0.03 and min(y + h, iy + ih) - max(y, iy) > 0.1:
                add('WARN', n, f'вертикальная линия перечёркивает текст «{t}»: разведите выноски или сгруппируйте вехи')
                break
            if h < 0.03 and iy + 0.04 < y < iy + ih - 0.04 and min(x + w, ix + iw) - max(x, ix) > 0.1:
                add('WARN', n, f'горизонтальная линия перечёркивает текст «{t}»')
                break

# серия: объекты в одной позиции на разных слайдах должны быть одного видимого размера
used = set()
for i, (n, cx, cy, _area) in enumerate(series):
    if i in used:
        continue
    grp = [j for j, (m, x2, y2, a2) in enumerate(series) if abs(x2 - cx) < 0.6 and abs(y2 - cy) < 0.6]
    if len({series[j][0] for j in grp}) >= 3:
        used |= set(grp)
        areas = [series[j][3] for j in grp]
        if max(areas) / max(min(areas), 1e-6) > 1.25:
            add('WARN', series[grp[0]][0], f'серия объектов разного видимого размера на слайдах {sorted({series[j][0] for j in grp})}: разброс x{max(areas) / min(areas):.1f}. Выровняйте через place3d()')

# повторы заголовков между слайдами
titles = defaultdict(list)
for n, slide in enumerate(prs.slides, 1):
    if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
        titles[slide.shapes.title.text_frame.text.strip()].append(n)
for t, ns in titles.items():
    if t and len(ns) > 1:
        add('WARN', ns[0], f'одинаковый заголовок на слайдах {ns}: «{t[:40]}»')

err = sum(i['level'] == 'ERR' for i in issues)
warn = sum(i['level'] == 'WARN' for i in issues)
for i in issues:
    print(f"[{i['level']}] слайд {i['slide']}: {i['msg']}")
print(f'\n{os.path.basename(a.pptx)}: {len(prs.slides)} слайдов, холст {SW:.3f}×{SH:.3f}″, ошибок {err}, предупреждений {warn}')
if a.json:
    json.dump({'errors': err, 'warnings': warn, 'issues': issues, 'slides': stats}, open(a.json, 'w'), ensure_ascii=False, indent=1)
sys.exit(1 if err or (a.strict and warn) else 0)

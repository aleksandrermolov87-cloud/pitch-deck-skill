# Полнота деки относительно брифа (reference/15-brief-coverage.md).
#   brief_coverage.py init  <src>                 заготовка out/coverage.json по requirements.json (для сабагента)
#   brief_coverage.py dump  <src> [--pptx файл]   текст каждого слайда и отдельно его заметки (вход для сабагента)
#   brief_coverage.py check <src> [--pptx файл] [--json отчёт.json]   карта «требование → слайд → элемент / заметки / нет», ошибки, out/coverage.md
# requirements.json (этап 1, до сборки): требования брифа к содержанию и месту.
# out/coverage.json (этап 6, независимый сабагент): где каждое требование выполнено и почему так.
# Скрипт не верит карте на слово: числа и ключевые слова требования он сам ищет в тексте слайдов и отдельно в заметках.
# Число, которое есть только в заметках, не засчитывается как показанное на слайде. Число, которое только рисует
# график (out/facts.json: shown = false), тоже: отчёт называет его «влияет на график, не подписано» и подсказывает opt.show.
# Число, которое читается, но без своего периода или единицы в той же подписи (shown_as = value: период только на оси),
# засчитывается как показанное с предупреждением «нарисовано и подписано числом, контекст отдельно».
# out/coverage.md: требования по группам (на слайдах / изменено ради читаемости / только в заметках / нет или сокращено /
# исключено) и карта по слайдам «слайд → требования». Код 1, если полнота не пройдена.
# Запуск: $SKILL/.venv/bin/python -I brief_coverage.py …
import argparse
import json
import os
import re
import sys

ap = argparse.ArgumentParser()
ap.add_argument('cmd', choices=['init', 'dump', 'check'])
ap.add_argument('src')
ap.add_argument('--pptx', help='готовый файл; по умолчанию ../<deck.json name>.pptx')
ap.add_argument('--req', help='requirements.json; по умолчанию <src>/requirements.json')
ap.add_argument('--map', help='coverage.json; по умолчанию <src>/out/coverage.json')
ap.add_argument('--json', help='check: итог в JSON — требований, выполнено, исключено, не выполнены (для eval_deck.py)')
a = ap.parse_args()

src = os.path.abspath(a.src)
req_p = a.req or os.path.join(src, 'requirements.json')
map_p = a.map or os.path.join(src, 'out', 'coverage.json')
if not os.path.isfile(req_p) and a.cmd != 'dump':
    sys.exit(f'нет {req_p}: требования брифа выписываются на этапе 1 (reference/15-brief-coverage.md)')
REQ = json.load(open(req_p, encoding='utf-8')) if os.path.isfile(req_p) else {'requirements': [{'id': '-', 'text': '-'}]}
R = REQ.get('requirements', [])
STATUSES = {'ok', 'partial', 'notes_only', 'missing', 'excluded'}
MUST = {'slide', 'notes', 'any'}

errors, warns = [], []
ids = [r.get('id') for r in R]
if not R:
    errors.append('requirements.json: нет ни одного требования')
for r in R:
    if not r.get('id') or ids.count(r['id']) > 1:
        errors.append(f"requirements.json: id «{r.get('id')}» пустой или повторяется")
    if not r.get('text'):
        errors.append(f"{r.get('id')}: нет текста требования")
    if r.get('must', 'slide') not in MUST:
        errors.append(f"{r.get('id')}: must = slide | notes | any")

if a.cmd == 'init':
    os.makedirs(os.path.dirname(map_p), exist_ok=True)
    if os.path.isfile(map_p):
        sys.exit(f'{map_p} уже есть: не перезаписываю')
    items = [{'id': r['id'], 'status': 'todo', 'slide': None, 'element': None, 'evidence': '', 'paraphrase': '', 'reason': ''} for r in R]
    json.dump({'by': 'сабагент-проверяющий', 'items': items}, open(map_p, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'заготовка {map_p}: {len(items)} требований')
    sys.exit(0)


# ── тексты слайдов и заметок из готового PPTX: то, что увидит зритель ──
def find_pptx():
    if a.pptx:
        return a.pptx
    deck = os.path.join(src, 'deck.json')
    name = json.load(open(deck, encoding='utf-8')).get('name') if os.path.isfile(deck) else None
    p = os.path.join(os.path.dirname(src), f'{name}.pptx') if name else ''
    return p if os.path.isfile(p) else None


def shape_texts(sh):
    out = []
    if sh.shape_type == 6:                                   # группа
        for s in sh.shapes:
            out += shape_texts(s)
    if getattr(sh, 'has_text_frame', False) and sh.has_text_frame:
        out.append(sh.text_frame.text)
    if getattr(sh, 'has_table', False) and sh.has_table:
        out += [c.text for row in sh.table.rows for c in row.cells]
    if getattr(sh, 'has_chart', False) and sh.has_chart:
        ch = sh.chart
        if ch.has_title:
            out.append(ch.chart_title.text_frame.text)
        for pl in ch.plots:
            out += [str(c) for c in pl.categories]
            for se in pl.series:
                out.append(se.name or '')
                out += [str(v) for v in se.values if v is not None]
    return out


pptx = find_pptx()
if not pptx or not os.path.isfile(pptx):
    sys.exit('нет PPTX: укажите --pptx <файл> (полнота проверяется по готовому файлу, а не по исходникам)')
from pptx import Presentation   # noqa: E402  (python-pptx есть в venv скилла)
P = Presentation(pptx)
SLIDES = []
for i, s in enumerate(P.slides, 1):
    body = '\n'.join(t for sh in s.shapes for t in shape_texts(sh) if t.strip())
    notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide and s.notes_slide.notes_text_frame else ''
    SLIDES.append({'n': i, 'body': body, 'notes': notes})
man_p = os.path.join(src, 'out', 'slides.json')
MAN = {s['n']: s for s in json.load(open(man_p, encoding='utf-8'))} if os.path.isfile(man_p) else {}
COVER = next((n for n, s in MAN.items() if s.get('archetype') == 'cover'), 1)
facts_p = os.path.join(src, 'out', 'facts.json')
FACTS = json.load(open(facts_p, encoding='utf-8')) if os.path.isfile(facts_p) else []
if a.cmd == 'dump':
    for s in SLIDES:
        ids_ = ', '.join(f"{t['id']}" for t in MAN.get(s['n'], {}).get('texts', []) if t.get('id'))
        print(f"\n═══ слайд {s['n']}" + (f" [{MAN[s['n']].get('archetype')}; компоненты: {'+'.join(MAN[s['n']].get('lib') or []) or '—'}]" if s['n'] in MAN else '')
              + (f"\n  id текстов: {ids_}" if ids_ else ''))
        print('НА СЛАЙДЕ:\n' + (s['body'].strip() or '—'))
        print('ЗАМЕТКИ (зритель их не видит):\n' + (s['notes'].strip() or '—'))
    sys.exit(0)


def norm(s):
    s = (s or '').replace('ё', 'е').replace('Ё', 'Е')
    s = re.sub(r'[   ]', ' ', s)
    s = re.sub(r'(?<=\d) (?=\d{3}\b)', '', s)             # 9 505 → 9505
    s = re.sub(r'(?<=\d)\.(?=\d)', ',', s)                 # 31.6 → 31,6
    s = re.sub(r'[«»"“”„]', '', s)
    return re.sub(r'\s+', ' ', s).strip().casefold()


def has(text, key):
    k = norm(key)
    if re.fullmatch(r'[\d,]+', k):                         # число: не часть другого числа
        return re.search(rf'(?<![\d,]){re.escape(k)}(?![\d]|,\d)', norm(text)) is not None
    return k in norm(text)


def alts(k):
    return k if isinstance(k, list) else [k]


# слова-инструкции брифа («показать», «вывод 2», «на обложке») не входят в сравнение смысла
INSTR = {'показ', 'указа', 'назва', 'отраз', 'вывод', 'слайд', 'облож', 'требо', 'должн', 'выводы'}


def stems(s):
    return {w[:5] for w in re.findall(r'\w+', norm(s)) if len(w) > 3 and not w.isdigit()} - INSTR


def paraphrased(r, ev):
    # цитата со слайда — не та же фраза, что в брифе: нужно пояснение. Числа требования в цитате снимают вопрос
    if r.get('numbers') and all(any(has(ev, k) for k in alts(n)) for n in r['numbers']):
        return False
    a, b = stems(r['text']), stems(ev)
    return len(a & b) < 0.5 * max(1, min(len(a), len(b)))


def targets(r):
    # где бриф требует показать: обложка, конкретный слайд или любой слайд деки
    w = r.get('where', 'any')
    if w == 'cover':
        return [COVER]
    if isinstance(w, str) and w.startswith('slide:'):
        return [int(w.split(':')[1])]
    return [s['n'] for s in SLIDES]


def where_label(r):
    w = r.get('where', 'any')
    return 'обложка' if w == 'cover' else f"слайд {w.split(':')[1]}" if str(w).startswith('slide:') else 'любой слайд'


def key_value(key):
    try:
        return float(norm(key).replace(' ', '').replace(',', '.'))
    except ValueError:
        return None


def ident(f):
    # идентичность факта из реестра: компонент, путь в данных, ряд, период
    return ' '.join(x for x in (f.get('lib'), f.get('path'), f.get('series') if f.get('series') and f.get('path') else None,
                                f.get('period')) if x) or f.get('label', '')


def graph_only(key, nums):
    # число требования рисует компонент на этих слайдах (lib), но не подписано (shown = false в out/facts.json)
    v = key_value(key)
    if v is None:
        return [], []
    fs = [f for f in FACTS if f.get('shown') is False and f.get('lib') and f['slide'] in nums and abs(abs(f['value']) - v) < 1e-9]
    return sorted({f['slide'] for f in fs}), fs


def value_only(key, nums):
    # число видно, но ни один его факт на этих слайдах не подписан вместе с периодом и единицей (shown_as = value)
    v = key_value(key)
    if v is None:
        return []
    fs = [f for f in FACTS if f.get('shown') and f['slide'] in nums and abs(abs(f['value']) - v) < 1e-9]
    return fs if fs and all(f.get('shown_as', 'label') == 'value' for f in fs) else []


def where_is(key, nums):
    vis = [s['n'] for s in SLIDES if s['n'] in nums and any(has(s['body'], k) for k in alts(key))]
    note = [s['n'] for s in SLIDES if any(has(s['notes'], k) for k in alts(key))]
    anyvis = [s['n'] for s in SLIDES if any(has(s['body'], k) for k in alts(key))]
    return vis, note, anyvis


MAP = {}
if os.path.isfile(map_p):
    for it in json.load(open(map_p, encoding='utf-8')).get('items', []):
        MAP[it.get('id')] = it
else:
    errors.append(f'нет {map_p}: карту строит независимый сабагент (brief_coverage.py init, затем промпт из reference/15-brief-coverage.md)')
for k in MAP:
    if k not in ids:
        warns.append(f'coverage.json: «{k}» нет в requirements.json')

rows = []
soft, nolabel = [], set()   # подписано числом без периода и единицы: предупреждение, на полноту не влияет
hits = {}            # слайд → требования, чьи числа или ключи на нём видны (карта по слайдам)
for r in R:
    rid, must, req = r.get('id'), r.get('must', 'slide'), r.get('required', True)
    mech = set()     # что нашла механика: notes (только в заметках), graph (только в графике), elsewhere, absent
    it = MAP.get(rid)
    sev = errors if req else warns
    errs0 = len(errors) if req else len(errors) + len(warns)
    st = (it or {}).get('status')
    place = '—'
    if not it or st not in STATUSES:
        sev.append(f'{rid}: нет статуса в карте покрытия (ok / partial / notes_only / missing / excluded)')
        st = st or 'нет'
    elif st == 'excluded':
        if not (it.get('reason') or '').strip():
            errors.append(f'{rid}: исключение без причины')
        place = 'исключено'
    elif st in ('missing', 'notes_only', 'partial'):
        why = (it.get('reason') or '').strip()
        sev.append(f"{rid}: {'нет в деке' if st == 'missing' else 'только в заметках' if st == 'notes_only' else 'выполнено частично'}"
                   f" — «{r['text'][:70]}»" + (f'; причина: {why}' if why else ''))
        place = {'missing': 'нет', 'notes_only': f"заметки слайда {it.get('slide')}", 'partial': f"слайд {it.get('slide')}, частично"}[st]
    else:   # ok: карта должна показать, где и чем
        n, el, ev = it.get('slide'), it.get('element'), (it.get('evidence') or '').strip()
        s = next((x for x in SLIDES if x['n'] == n), None)
        if not s:
            errors.append(f'{rid}: в карте нет номера слайда или такого слайда нет ({n})')
        elif must != 'notes' and r.get('where') in ('cover',) and n != COVER:
            sev.append(f'{rid}: требование про обложку, а карта указывает слайд {n}')
        if s and not el:
            errors.append(f'{rid}: не указан элемент слайда (id текста, lib:<компонент>, visual:<что> или notes)')
        visual = isinstance(el, str) and el.startswith(('lib:', 'visual:'))
        if s and ev and not visual:
            if must != 'notes' and not has(s['body'], ev):
                if has(s['notes'], ev):
                    sev.append(f'{rid}: цитата есть только в заметках слайда {n}, а требование — показать на слайде')
                else:
                    errors.append(f'{rid}: цитаты «{ev[:50]}» нет в тексте слайда {n}')
            if paraphrased(r, ev) and not (it.get('paraphrase') or '').strip():
                errors.append(f'{rid}: формулировка на слайде отличается от брифа, нужно пояснение paraphrase (почему это то же утверждение)')
        if s and not ev:
            errors.append(f'{rid}: нет evidence — текста со слайда или описания визуального утверждения')
        if s and visual:
            lib = el.split(':', 1)[1]
            if el.startswith('lib:') and MAN and lib not in (MAN.get(n, {}).get('lib') or []):
                errors.append(f'{rid}: на слайде {n} нет компонента {lib}')
            if not (r.get('numbers') or r.get('keys')):   # числа и ключи ниже проверяются сами, без них смысл видит только сабагент
                warns.append(f'{rid}: визуальное утверждение на слайде {n} без чисел и ключей — смысл подтверждает только сабагент по PNG')
        place = f"слайд {n}: {el}" if s else '—'
    # механика поверх карты: числа и ключевые слова требования ищем сами
    if st != 'excluded':
        tgt = targets(r)
        for kind, keys in (('число', r.get('numbers', [])), ('ключ', r.get('keys', []))):
            for k in keys:
                vis, note, anyvis = where_is(k, tgt)
                for n_ in vis:
                    hits.setdefault(n_, set()).add(rid)
                label = ' / '.join(alts(k))
                if must == 'notes':
                    if not note:
                        sev.append(f'{rid}: {kind} «{label}» нет в заметках')
                        mech.add('absent')
                elif vis and kind == 'число' and value_only(alts(k)[0], vis):
                    fs = value_only(alts(k)[0], vis)
                    soft.append(f"{rid}: число «{label}» на слайде {sorted({f['slide'] for f in fs})} подписано без периода или единицы "
                                f"в той же подписи ({ident(fs[0])}): нарисовано и читается, контекст только на оси или в подзаголовке")
                    nolabel.add(rid)
                elif not vis:
                    g, gf = graph_only(alts(k)[0], tgt) if kind == 'число' else ([], [])
                    if g:
                        sev.append(f'{rid}: число «{label}» влияет на график слайда {g} ({ident(gf[0])}), но его нельзя прочитать: подпишите его '
                                   f'(opt.show у компонента, подпись или текст рядом с графиком)')
                        mech.add('graph')
                    elif anyvis:
                        sev.append(f'{rid}: {kind} «{label}» есть на слайдах {anyvis}, но не там, где требует бриф ({where_label(r)})')
                        mech.add('elsewhere')
                    elif note:
                        sev.append(f'{rid}: {kind} «{label}» есть только в заметках (слайды {note}), на слайде его нет')
                        mech.add('notes')
                    else:
                        sev.append(f'{rid}: {kind} «{label}» нет ни на слайдах, ни в заметках')
                        mech.add('absent')
                    if st == 'ok':
                        place += ' ⚠'
    bad = (len(errors) if req else len(errors) + len(warns)) > errs0
    if it and isinstance(it.get('slide'), int) and st in ('ok', 'partial'):
        hits.setdefault(it['slide'], set()).add(rid)
    # группа для отчёта: что зритель увидит и что нет
    if st == 'excluded':
        bucket = 'excluded'
    elif st == 'notes_only' or (mech == {'notes'}):
        bucket = 'notes'
    elif st in ('missing', 'partial', 'нет') or mech or bad:
        bucket = 'lost'
    elif (it or {}).get('paraphrase'):
        bucket = 'edited'
    else:
        bucket = 'slide'
    rows.append((rid, r['text'], st, place, (it or {}).get('paraphrase') or (it or {}).get('reason') or '', bad, bucket,
                 sorted(mech | ({'nolabel'} if rid in nolabel else set()))))
warns += soft

# ── out/coverage.md: карта для REPORT.md ──
GROUPS = [('slide', 'На слайдах'), ('edited', 'На слайдах, формулировка изменена ради читаемости'), ('notes', 'Только в заметках докладчика'),
          ('lost', 'Нет на слайдах или сокращено'), ('excluded', 'Исключено с причиной')]
MECH = {'graph': 'число только в графике, без подписи', 'nolabel': 'число подписано без периода или единицы рядом', 'notes': 'только в заметках', 'elsewhere': 'не на том слайде', 'absent': 'нет в деке'}
ok = sum(1 for x in rows if x[6] in ('slide', 'edited'))
exc = sum(1 for x in rows if x[6] == 'excluded')
md = ['## Полнота относительно брифа', '',
      f'Требований {len(rows)}: на слайдах {ok} (из них с изменённой формулировкой {sum(1 for x in rows if x[6] == "edited")}), '
      f'только в заметках {sum(1 for x in rows if x[6] == "notes")}, нет или сокращено {sum(1 for x in rows if x[6] == "lost")}, '
      f'исключено с причиной {exc}. Ошибок {len(errors)}, предупреждений {len(warns)}.', '']
for key, title in GROUPS:
    grp = [x for x in rows if x[6] == key]
    if not grp:
        continue
    md += [f'### {title} ({len(grp)})', '', '| # | Требование | Статус | Где в деке | Пояснение |', '|---|---|---|---|---|']
    for rid, text, st, place, note, bad, _, mech in grp:
        why = '; '.join([note] + [MECH[m] for m in mech if m in MECH]).strip('; ')
        md.append(f"| {rid} | {text} | {st}{' ⚠' if bad else ''} | {place} | {why} |")
    md.append('')
BAD = {x[0] for x in rows if x[6] in ('notes', 'lost')}
md += ['### Карта по слайдам', '', 'Требование с ⚠ выполнено на слайде не полностью или только в заметках.', '', '| Слайд | Заголовок | Компоненты | Требования брифа |', '|---|---|---|---|']
for sl in SLIDES:
    m = MAN.get(sl['n'], {})
    rq = ', '.join(x + (' ⚠' if x in BAD else '') for x in sorted(hits.get(sl['n'], set()), key=lambda x: (len(x), x))) or ('— (обложка или связка)' if sl['n'] == COVER else '— не закрывает ни одного требования')
    md.append(f"| {sl['n']} | {(m.get('title') or sl['body'].splitlines()[0] if sl['body'] else '')[:70]} | {'+'.join(m.get('lib') or []) or m.get('custom') or '—'} | {rq} |")
md.append('')
os.makedirs(os.path.join(src, 'out'), exist_ok=True)
open(os.path.join(src, 'out', 'coverage.md'), 'w', encoding='utf-8').write('\n'.join(md))
for e in errors:
    print('[ERR]', e)
for w in warns:
    print('[WARN]', w)
print(f"\nтребований {len(rows)}, выполнено {ok}, исключено {exc}, ошибок {len(errors)}, предупреждений {len(warns)} — "
      f"полнота {'пройдена' if not errors else 'НЕ пройдена'}; карта: {os.path.join(src, 'out', 'coverage.md')}")
if a.json:
    json.dump({'total': len(rows), 'ok': ok, 'excluded': exc, 'failed': [x[0] for x in rows if x[6] in ('notes', 'lost')],
               'errors': len(errors), 'warnings': len(warns)}, open(a.json, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
sys.exit(1 if errors else 0)

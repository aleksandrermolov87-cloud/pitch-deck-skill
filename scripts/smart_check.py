# Проверка тезисов и буллетов по SMART (reference/01-brief-and-copy.md): есть ли число с единицей (M) и срок (T),
# не размыт ли объект (S). Слайды плана, роадмапа, целей и статуса определяются по словам в заголовке или задаются --slides.
# Запуск: $SKILL/.venv/bin/python -I smart_check.py <deck.pptx> [--slides 1,2,13] [--min 0.6] [--json report.json]
# Код 1, если доля SMART-пунктов на проверяемых слайдах ниже --min.
import argparse
import json
import re
import sys

from pptx import Presentation

ap = argparse.ArgumentParser()
ap.add_argument('pptx')
ap.add_argument('--slides', default='', help='номера слайдов через запятую; по умолчанию — по ключевым словам заголовка')
ap.add_argument('--min', type=float, default=0.6, help='минимальная доля пунктов с числом и сроком')
ap.add_argument('--json')
a = ap.parse_args()

PLAN_WORDS = r'роадмап|roadmap|план|цел|пилот|запуск|релиз|направлен|итог|результат|эффект|просим|утверд|статус|планируем|сделано|следующ'
MEASURE = r'\d[\d\s ]*(?:[.,]\d+)?\s*(?:%|п\.п\.|млн|млрд|тыс|шт|₽|руб|полис|клиент|продукт|пилот|эпик|направлен|канал|раз|месяц|недел|дн|x\d)|x\d+|\b\d+\s*(?:из|/)\s*\d+'
TIME = (r'\b(?:I{1,3}V?|IV)\s*кв|\b[1-4]\s*кв|квартал|\b20\d\d\b|\bМ\d{1,2}\b|январ|феврал|март|апрел|ма[йя]|июн|июл|август|сентябр|'
        r'октябр|ноябр|декабр|\bк\s+\d|\bдо\s+\d|\bс\s+\d|недел|месяц|за\s+год|\d{1,2}\.\d{1,2}(?:\.\d{2,4})?')
VAGUE = r'^(?:развити[ея]|улучшени[ея]|оптимизаци[яи]|повышени[ея]|работа\s+над|проработк[аи]|поддержк[аи]|внедрени[ея]|совершенствовани[ея])\b'

prs = Presentation(a.pptx)
pick = {int(x) for x in a.slides.split(',') if x.strip()}
rows, checked = [], 0
for n, slide in enumerate(prs.slides, 1):
    if slide._element.get('show') == '0':
        continue
    texts = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    if not texts:
        continue
    title = max(texts, key=lambda sh: max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0)).text_frame.text
    if pick and n not in pick:
        continue
    if not pick and not re.search(PLAN_WORDS, title, re.I):
        continue
    # срок может задаваться осью (кварталы, месяцы в подписях шкалы): тогда пункты строк получают T от оси
    axis = sum(1 for sh in texts for p in sh.text_frame.paragraphs
               if len(p.text.strip()) <= 12 and re.search(TIME + r'|^(?:I{1,3}V?|IV)$|^Янв|^Фев|^Мар|^Апр|^Май|^Июн|^Июл|^Авг|^Сен|^Окт|^Ноя|^Дек', p.text.strip(), re.I))
    on_axis = axis >= 4
    for sh in texts:
        if sh.top is not None and sh.top / 914400 < 1.85:      # шапка: заголовок, плашка, подзаголовок — не пункты
            continue
        for p in sh.text_frame.paragraphs:
            t = p.text.strip()
            # пункт: 5–20 слов (короче — подпись продукта или канала), не чисто число
            words = t.split()
            if not (5 <= len(words) <= 20) or re.fullmatch(r'[\d\s.,%₽x–-]+', t) or t == title.strip():
                continue
            m = bool(re.search(MEASURE, t, re.I))
            tm = bool(re.search(TIME, t, re.I)) or on_axis
            vague = bool(re.search(VAGUE, t, re.I))
            checked += 1
            rows.append({'slide': n, 'text': t, 'M': m, 'T': tm, 'T_axis': on_axis and not re.search(TIME, t, re.I), 'vague_S': vague})

smart = [r for r in rows if r['M'] and r['T'] and not r['vague_S']]
share = len(smart) / checked if checked else 1.0
for r in rows:
    miss = [k for k, ok in (('M число', r['M']), ('T срок', r['T'])) if not ok] + (['S размыто'] if r['vague_S'] else [])
    if miss:
        print(f"[WARN] слайд {r['slide']}: «{r['text'][:70]}» — нет: {', '.join(miss)}")
print(f'\nпроверено пунктов {checked}, SMART (число + срок + конкретный объект) {len(smart)}, доля {share:.0%}, порог {a.min:.0%}')
print('Не выдумывай число или срок ради SMART: если их нет в источнике, вынеси вопрос в отчёт.')
if a.json:
    json.dump({'checked': checked, 'smart': len(smart), 'share': share, 'rows': rows}, open(a.json, 'w'), ensure_ascii=False, indent=1)
sys.exit(1 if share < a.min else 0)

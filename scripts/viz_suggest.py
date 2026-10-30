# Детектор визуализаций (под-скилл «Схемы», reference/19-schemes.md): ищет в тексте то, что читается лучше картинкой,
# и предлагает компонент библиотеки. Для схемы связей собирает черновик данных schemeMap: участники, пункты, связи, хаб.
#   viz_suggest.py brief <файл…>          вводные: .md .txt .docx .pptx .pdf (pdf — через pdftotext, если есть)
#   viz_suggest.py deck  <src-проекта>    собранная дека (out/slides.json): текстовые слайды, которые просятся в схему,
#                                          и блоковые слайды из PPTX (ряд или сетка одинаковых карточек) — их заменить обязательно
#   --json                                 список предложений в JSON (вход для агента)
#   --emit <папка>                         черновики данных: scheme-<k>.json для каждой предложенной схемы связей
# Предложения — повод спросить пользователя, а не решение: агент показывает их одним вопросом (SKILL.md, этапы 0 и 2).
# Черновик схемы — каркас: подписи связей (tag) и короткие формулировки пишет агент по тексту, числа не выдумываются.
# Запуск: $SKILL/.venv/bin/python -I viz_suggest.py …
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile

ap = argparse.ArgumentParser()
ap.add_argument('cmd', choices=['brief', 'deck'])
ap.add_argument('paths', nargs='+')
ap.add_argument('--json', action='store_true')
ap.add_argument('--emit')
ap.add_argument('--pptx', help='deck: готовый файл для поиска блоковых слайдов; по умолчанию ../<deck.json name>.pptx')
a = ap.parse_args()


# ───────── чтение ─────────
def read(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.md', '.txt', ''):
        return open(path, encoding='utf-8', errors='replace').read()
    if ext == '.docx':
        from defusedxml import ElementTree as ET
        with zipfile.ZipFile(path) as z:
            root = ET.fromstring(z.read('word/document.xml'))
        ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        return '\n'.join(''.join(t.text or '' for t in p.iter(ns + 't')) for p in root.iter(ns + 'p'))
    if ext == '.pptx':
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(path).slides, 1):
            out.append(f'# Слайд {i}')
            out += [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        return '\n'.join(out)
    if ext == '.pdf':
        if not shutil.which('pdftotext'):
            sys.exit(f'{path}: для PDF нужен pdftotext (poppler: zsh setup.sh --system)')
        return subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, text=True, check=True).stdout
    sys.exit(f'{path}: формат не поддерживается (.md .txt .docx .pptx .pdf)')


# ───────── признаки ─────────
LINK = re.compile(r'взаимодейств|обратн\w* связ|\bсвяз[ьи]\b|переда[её]т|согласу|синхрониз|координ|единое окно|\bsla\b|договор|сверк|'
                  r'коммуникац|отчитыва|обмен|интеграц|канал|партн[её]р|эскалац|бэклог|подключа|предоставля|получа[ею]т', re.I)
STOP = re.compile(r'^(итог|цель|цели|срок|шаг|этап|риск|вопрос|задач|результат|пример|важно|примечание|комментар|источник|'
                  r'вывод|контекст|проблем|решени|метрик|kpi|бюджет|план|статус|слайд|тема|аудитор|дата|время)', re.I)
KINDS = [
    # (компонент, что показать, вопрос зрителя, regex, сколько совпадений нужно). Схемы — reference/19-schemes.md
    ('hubSpoke', 'хаб и участники вокруг', 'кто вокруг центра и что каждый с ним делает?', r'единое окно|центр\w* (компетенц|взаимодейств)|координир|вокруг (продукта|платформы)', 1),
    ('contextDiagram', 'контекст системы', 'с кем и чем обменивается система?', r'внешн\w* систем|пользовател\w* системы|интегрир\w* с|контур банка', 2),
    ('contourMap', 'контур: внутри и снаружи', 'что делаем сами, что отдаём наружу?', r'внутри (банка|команды|контура)|снаружи|аутсорс|на стороне партн[её]ра|периметр', 2),
    ('archFlows', 'архитектура с потоками', 'как запрос проходит через систему?', r'\bслой|уровень систем|архитектур|фронт|бэкенд|backend|\bapi\b|шина|хранилищ', 3),
    ('radialHub', 'радиальная схема', 'кто вокруг центра?', r'вокруг|в центре|ядро', 2),
    ('cycleLoop', 'цикл', 'как повторяется процесс?', r'\bцикл|повторя|по кругу|петл\w* обратной связи|итерац', 2),
    ('flywheel', 'маховик', 'что разгоняет рост?', r'маховик|flywheel|самоусилива|ускоря|разгоня', 1),
    ('orbitRings', 'экосистема кольцами', 'кто в ядре, кто в ближнем и дальнем круге?', r'экосистем|ближн\w* круг|дальн\w* круг|\bядро', 2),
    ('partnerTriangle', 'треугольник сторон', 'как связаны три стороны?', r'банк\w*.{0,60}страхов\w*.{0,60}клиент|клиент\w*.{0,60}банк\w*.{0,60}страхов|трёхсторон|трехсторон', 1),
    ('vennOverlap', 'пересечение зон', 'где зоны пересекаются?', r'пересечени|общ\w* зон|на стыке|совместн\w* ответствен', 2),
    ('decisionTree', 'дерево решений', 'что делать в каждом случае?', r'\bесли\b.{0,80}\bто\b|в случае|иначе|да\s*/\s*нет', 2),
    ('driverTree', 'дерево драйверов', 'из чего складывается цель?', r'драйвер|рычаг|декомпозиц|из чего складыва', 2),
    ('governanceTree', 'уровни управления', 'кто что решает и как часто?', r'комитет|правлени|\bсовет\b|полномоч|кворум', 2),
    ('mindMap', 'карта темы', 'какие у темы направления?', r'направлени[яй] работы|аспект|ветк', 2),
    ('swimlaneFlow', 'процесс по дорожкам', 'кто что делает на каждом шаге?', r'переда[её]т|передача|получает от|после согласования|возвращает на доработку', 3),
    ('serviceBlueprint', 'схема сервиса', 'что видит клиент и что за кадром?', r'клиент видит|точк\w* контакта|фронт-?офис|бэк-?офис|за кадром', 2),
    ('sequenceDiagram', 'последовательность вызовов', 'в каком порядке идут запросы?', r'запрос|ответ\b|отправля|вызыва|callback|вебхук|webhook', 4),
    ('handoffTimeline', 'передачи во времени', 'кто кому и когда передаёт работу?', r'переда[её]т.{0,40}(в|к) (недел|месяц|квартал|срок)|к (концу|началу) (месяца|квартала)|после (запуска|пилота)', 2),
    ('escalationLadder', 'лестница эскалации', 'куда идёт вопрос, если не решён?', r'эскалац|первая линия|вторая линия|третья линия|\bL[1-3]\b', 2),
    ('stateMachine', 'статусная модель', 'через какие статусы проходит объект?', r'статус|переходит в|жизненн\w* цикл|расторж|аннулир', 3),
    ('dependencyGraph', 'граф зависимостей', 'что от чего зависит?', r'зависит от|блокирует|предпосылк|после (запуска|готовности)|\bволн', 2),
    ('dataFlow', 'потоки данных', 'какие данные куда идут?', r'данные|выгрузк|загрузк|интеграц|поток', 3),
    ('impactChain', 'цепочка влияния', 'как инициатива доходит до эффекта?', r'влияет на|вед[её]т к|приводит к|эффект|метрик', 3),
    ('busIntegration', 'интеграционная шина', 'как системы подключены друг к другу?', r'шина|esb|брокер|kafka|интеграционн', 1),
    ('bridgeMapping', 'сопоставление было → стало', 'что во что переходит?', r'было\s*[→>-]|маппинг|сопоставлени|переход\w* в новую|as is.{0,20}to be', 1),
    ('valueExchange', 'обмен ценностью', 'что стороны дают друг другу?', r'взамен|получает|да[её]т|обмен', 3),
    ('platformTwoSided', 'двусторонняя платформа', 'что получают обе стороны платформы?', r'двусторон|маркетплейс|платформ\w* для|продавц\w*.{0,40}покупател', 1),
    ('matrixOrg', 'матричная структура', 'как направления пересекаются с функциями?', r'матричн|двойное подчинение|функци\w* и направлени|кросс-?функц', 1),
    ('causeEffect', 'диаграмма причин', 'почему возникает проблема?', r'причин|из-за|следстви|корн\w* проблем', 3),
    ('orgChart', 'оргструктура', 'кто кому подчиняется?', r'подчин|руководител|в подчинении|\bотдел|штат|ваканси', 2),
    ('deptTree', 'дерево департамента', 'кто кому подчиняется в департаменте и сколько людей в звене?', r'департамент|управлени[ея] \w+|\bотдел\w*|штат\w*|ваканси', 3),
    ('orgOutline', 'оргструктура с колонками', 'сколько людей и вакансий в каждом звене?', r'штатн\w* расписани|численност|штат\w*|ваканси|\bfte\b', 3),
    ('spanOfControl', 'норма управляемости', 'у кого слишком много подчинённых?', r'норм\w* управляемост|span of control|прям\w* подчинённ|подчинённых у', 1),
    ('dualReporting', 'двойное подчинение', 'кому подчиняются административно и функционально?', r'функционально подчин|административно подчин|функциональн\w* руковод|пунктирн\w* подчин|dotted line', 1),
    ('orgNest', 'карта подразделений по численности', 'из чего состоит блок и какого размера части?', r'численност|штат\w*|человек в|сотрудник\w* в', 3),
    ('teamTopology', 'топология команд', 'как команды работают вокруг потоков?', r'потоков\w* команд|stream-aligned|платформенн\w* команд|enabling|team topolog', 1),
    ('tribeChapters', 'трайбы и чаптеры', 'как пересекаются команды и компетенции?', r'\bтрайб|\bсквад|\bчаптер|гильди|tribe|squad|chapter', 2),
    ('orgTransition', 'было → стало в оргструктуре', 'что меняется при реорганизации?', r'реорганизац|объединя\w* (отдел|управлен|подразделен)|упраздн|выделя\w* в отдельн|новая структура', 1),
    ('regionNetwork', 'региональная сеть', 'как устроена сеть центр — регионы — точки?', r'регион\w* дирекц|филиал|головн\w* офис|точ\w* продаж|региональн\w* сет', 2),
    ('gradePyramid', 'пирамида уровней', 'сколько людей на каждом уровне?', r'\bN-[1-4]\b|грейд|уровн\w* управлен|топ-менедж|линейн\w* руковод', 2),
    ('relationshipWeb', 'сеть отношений', 'кто с кем работает теснее всего?', r'тесно работа|чаще всего взаимодейств|ключев\w* стейкхолдер|сеть контактов|стейкхолдер', 2),
    ('interactionMatrix', 'матрица взаимодействий', 'кто от кого зависит и насколько плотно?', r'зависит от|обращается к|запрашивает у|взаимозависим', 3),
    ('arcHandoffs', 'дуги передач', 'где работа переходит и возвращается?', r'переда[её]т в|возвращает|на доработку|передача между|хэндофф|handoff', 3),
    ('raciMatrix', 'матрица ответственности', 'кто исполняет, кто утверждает?', r'\braci\b|утвержда|согласует|информиру|исполнител', 3),
    ('processSteps', 'процесс по шагам', 'как проходит процесс?', r'(?m)^\s*(шаг|этап)\s*\d|^\s*\d+[.)]\s+\S|затем|после этого|→|->', 3),
    ('ganttLanes', 'сроки по дорожкам', 'кто что делает параллельно и когда?', r'\b(?:q[1-4]|[1-4]\s*кв\.?|м\d{1,2}|январ\w*|феврал\w*|март\w*|апрел\w*|июн\w*|июл\w*|август\w*|сентябр\w*|октябр\w*|ноябр\w*|декабр\w*|20\d\d)\b', 4),
    ('comparison', 'сравнение вариантов', 'что выбрать?', r'\bvs\.?\b|\bпротив\b|вариант [абв123]|альтернатив|сравнени', 2),
    ('funnel', 'воронка', 'где теряем клиентов?', r'воронк|конверси', 2),
    ('matrix2x2', 'позиционирование', 'где мы среди игроков?', r'конкурент|позиционир|\bигрок', 2),
]


def abbr(name):
    m = re.search(r'\(([A-ZА-ЯЁ]{2,6})\)', name)
    if m:
        return m.group(1)
    ws = [w for w in re.findall(r'[A-Za-zА-Яа-яЁё]+', name) if len(w) > 2]
    return ''.join(w[0].upper() for w in ws) if len(ws) >= 2 else None


def norm(name):
    return re.sub(r'\s*\([A-ZА-ЯЁ]{2,6}\)\s*', ' ', name.replace('**', '')).strip(' .')


def actors(text):
    """Участники: заголовки ##–#### с пунктами под ними и лиды строк «Имя: …» / «Имя — …» вне таких разделов.
    Заголовок первого уровня — название документа, не участник. «Команда (КС)» и «Команда» — один участник с сокращением КС."""
    found, alias, cur, under_head = {}, {}, None, False
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        h = re.match(r'^(#{1,4})\s+(.+?)\s*$', line)
        bullet = re.match(r'^[-*•]\s+(.+)$', line)
        lead = re.match(r'^(?:[-*•]\s*)?(?:\*\*)?([A-ZА-ЯЁ][^:—–\n]{1,48}?)(?:\*\*)?\s*(?::|\s[—–-]\s)\s*(.+)$', line)
        if h:
            name = norm(h.group(2))
            under_head = len(h.group(1)) > 1 and len(name.split()) <= 5 and not STOP.match(name)
            cur = name if under_head else None
            if cur:
                found.setdefault(cur, [])
                m = re.search(r'\(([A-ZА-ЯЁ]{2,6})\)', h.group(2))
                if m:
                    alias[cur] = m.group(1)
        elif bullet and cur and under_head:
            found[cur].append(bullet.group(1).replace('**', '').strip(' .'))
        elif lead and len(lead.group(1).split()) <= 5 and not STOP.match(lead.group(1).strip()):
            name = norm(lead.group(1))
            m = re.search(r'\(([A-ZА-ЯЁ]{2,6})\)', lead.group(1))
            if m:
                alias[name] = m.group(1)
            cur, under_head = name, False
            found.setdefault(cur, []).extend([p.strip(' .') for p in re.split(r';\s*', lead.group(2)) if p.strip(' .')])
        elif bullet and cur:
            found[cur].append(bullet.group(1).replace('**', '').strip(' .'))
        elif not under_head:
            cur = None
    return {k: v for k, v in found.items() if k}, alias


def mentions(line, names):
    hit = []
    for n, al in names.items():
        keys = [n] + ([al] if al else [])
        if any(re.search(r'(?<![\wА-Яа-яЁё])' + re.escape(k) + r'(?![\wА-Яа-яЁё])', line, re.I if k == n else 0) for k in keys):
            hit.append(n)
    return hit


def scheme(text):
    acts, alias = actors(text)
    names = {n: alias.get(n) or abbr(n) for n in acts}
    if len(names) < 3:
        return None
    edges, cues = {}, set()
    lead_of = None
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        lm = re.match(r'^(?:[-*•]\s*)?(?:\*\*)?([A-ZА-ЯЁ][^:—–\n]{1,48}?)(?:\*\*)?\s*(?::|\s[—–-]\s)', line)
        hm = re.match(r'^#{1,4}\s+(.+?)\s*$', line)
        if hm:
            lead_of = norm(hm.group(1)) if norm(hm.group(1)) in names else None
        elif lm and norm(lm.group(1)) in names:
            lead_of = norm(lm.group(1))
        elif not line.startswith(('-', '*', '•')):
            lead_of = None
        hit = mentions(line, names)
        if lead_of and lead_of not in hit:
            hit = [lead_of] + hit
        if len(hit) < 2:
            continue
        for m in LINK.finditer(line):
            cues.add(m.group(0).lower())
        body = re.sub(r'^[-*•#\s]+', '', line).replace('**', '')
        for o in hit[1:]:
            k = tuple(sorted((hit[0], o)))
            if k[0] != k[1] and k not in edges:
                edges[k] = body[:70] + ('…' if len(body) > 70 else '')
    if len(edges) < 2 and len(cues) < 2:
        return None
    deg = {n: 0 for n in names}
    for f, t in edges:
        deg[f] += 1
        deg[t] += 1
    hub = next((n for n in names if re.search(r'единое окно|координ|центр', ' '.join(acts[n]), re.I)), None) or max(deg, key=deg.get)
    ids = {n: f'n{i + 1}' for i, n in enumerate(names)}
    data = {
        'nodes': [{'id': ids[n], 'name': n, **({'kind': 'hub'} if n == hub else {}), 'items': acts[n][:4]} for n in names][:9],
        'edges': [{'from': ids[f], 'to': ids[t], 'tag': '', 'text': txt} for (f, t), txt in edges.items() if f in ids and t in ids],
    }
    return {'component': 'schemeMap', 'what': 'схема связей', 'question': 'кто с кем и о чём взаимодействует?',
            'why': f'{len(names)} участника(ов), {len(edges)} связ(ей), хаб «{hub}»' + (f'; признаки: {", ".join(sorted(cues)[:5])}' if cues else ''),
            'score': len(names) + 2 * len(edges) + len(cues), 'data': data}


def detect(text):
    out = []
    s = scheme(text)
    if s:
        out.append(s)
    for comp, what, q, rx, need in KINDS:
        hits = re.findall(rx, text, re.I)
        if len(hits) >= need:
            sample = sorted({(h if isinstance(h, str) else h[0]).strip().lower() for h in hits if (h if isinstance(h, str) else h[0]).strip()})[:4]
            out.append({'component': comp, 'what': what, 'question': q, 'why': f'{len(hits)} признак(ов): {", ".join(sample)}', 'score': len(hits)})
    return sorted(out, key=lambda r: -r['score'])


# ───────── режимы ─────────
res = []
if a.cmd == 'brief':
    text = '\n'.join(read(p) for p in a.paths)
    for r in detect(text):
        res.append({'slide': None, **r})
else:
    man = os.path.join(a.paths[0], 'out', 'slides.json')
    if not os.path.isfile(man):
        sys.exit(f'нет {man}: сначала node build.js')
    blocks = {}
    deck_json = os.path.join(a.paths[0], 'deck.json')
    pp = a.pptx or (os.path.join(os.path.dirname(os.path.abspath(a.paths[0])), json.load(open(deck_json, encoding='utf-8')).get('name', '') + '.pptx') if os.path.isfile(deck_json) else '')
    if pp and os.path.isfile(pp):
        from pptx import Presentation
        EMU = 914400
        for i, sl in enumerate(Presentation(pp).slides, 1):
            cards = []
            shapes = list(sl.shapes)
            for sh in shapes:
                if sh.shape_type == 1 and sh.width and sh.height and not (sh.has_text_frame and sh.text_frame.text.strip()):
                    x, y, w, h = sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU
                    inner = [t for t in shapes if t.has_text_frame and t.text_frame.text.strip() and t.left is not None
                             and t.left / EMU >= x - 0.01 and t.top / EMU >= y - 0.01 and (t.left + t.width) / EMU <= x + w + 0.01 and (t.top + t.height) / EMU <= y + h + 0.01]
                    if inner and w < 6 and sum(len(t.text_frame.text) for t in inner) > 14:
                        cards.append((x, y, w, h))
            same = max(([d for d in cards if abs(d[2] - c[2]) <= 0.08 * c[2] and abs(d[3] - c[3]) <= 0.08 * c[3]] for c in cards), key=len, default=[])
            if len(same) >= 3:
                blocks[i] = len(same)
    VISUAL = re.compile(r'^(scheme|hub|org|raci|process|gantt|timeline|systems|value|matrix|funnel|comparison|before)', re.I)
    for s in json.load(open(man, encoding='utf-8')):
        if s.get('hidden'):
            continue
        libs = s.get('lib') or []
        block = blocks.get(s['n']) if not (s.get('archetype') or '').startswith('scheme-') else None
        if any(VISUAL.match(x) for x in libs) and not block:
            continue                       # слайд уже нарисован схемой или графиком
        body = '\n'.join(''.join(t['runs']) for t in s.get('texts', []) if not re.search(r'-(title|sub|src)$', t.get('id') or '') and (t.get('lib') or block))
        if not block:
            body = body if body.strip() else ''
        if len(body.split()) < 20 and not block:
            continue
        found = detect(body)[:2]
        if block and not found:
            found = [{'component': 'schemeMap', 'what': 'схема связей или раскладка без контейнеров', 'question': 'как связаны эти карточки?',
                      'why': 'есть связи между карточками — schemeMap / hubSpoke; это шаги — processSteps / swimlaneFlow; это тезисы — principles / keyTakeaways', 'score': 0}]
        for r in found:
            res.append({'slide': s['n'], 'title': s.get('title', ''), 'now': ', '.join(libs) or s.get('archetype', ''),
                        **({'block': block} if block else {}), **r})

if a.emit:
    os.makedirs(a.emit, exist_ok=True)
    k = 0
    for r in res:
        if r['component'] == 'schemeMap':
            k += 1
            r['draft'] = os.path.join(a.emit, f'scheme-{k}.json')
            json.dump(r['data'], open(r['draft'], 'w', encoding='utf-8'), ensure_ascii=False, indent=2)

if a.json:
    print(json.dumps(res, ensure_ascii=False, indent=2))
    sys.exit(0)
if not res:
    print('визуализаций не предложено: в тексте нет участников со связями, шагов, сроков, слоёв или сравнений')
    sys.exit(0)
print(f'Предлагаю визуализации ({len(res)}):')
for i, r in enumerate(res, 1):
    where = (f'слайд {r["slide"]} (сейчас: {r["now"] or "текст"}' + (f'; БЛОКОВЫЙ: {r["block"]} одинаковых карточки — заменить обязательно' if r.get('block') else '') + ') → ') if r['slide'] else ''
    print(f'  {i}. {where}{r["what"]} — {r["component"]}: «{r["question"]}»')
    print(f'     {r["why"]}')
    if r.get('draft'):
        print(f'     черновик данных: {r["draft"]}')
